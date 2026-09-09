"""公共工具:范围解析、值序列化、工作簿加载/原子保存。"""

from __future__ import annotations

import os
import re
import tempfile
from datetime import date, datetime, time
from time import sleep as _sleep
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from ._atomic import replace_with_retry
from .errors import CliError

#: openpyxl 支持的扩展名;.xls(旧二进制格式)不支持
SUPPORTED_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
MAX_ROWS = 1_048_576
MAX_COLS = 16_384

# ---------------------------------------------------------------------------
# 列号 <-> 字母
# ---------------------------------------------------------------------------

def idx_to_col(n: int) -> str:
    """1 -> A, 27 -> AA"""
    letters = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def col_to_idx(text: str) -> int:
    """A -> 1, AA -> 27"""
    n = 0
    for ch in text.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


# ---------------------------------------------------------------------------
# 范围解析:支持 B2 / A1:B3 / A:C(整列) / 1:3(整行),可带 $ 前缀
# ---------------------------------------------------------------------------

def parse_range(text: str, ws: Optional[Worksheet] = None) -> tuple[int, int, int, int]:
    """解析范围文本,返回 (min_row, min_col, max_row, max_col) 1-based。

    整列/整行依赖 ws 的实际尺寸展开;ws 为 None 时整列/整行报错。
    """
    raw = text.strip().replace("$", "")
    if not raw:
        raise CliError("bad_range", f"范围不能为空,应为 B2 / A1:B3 / A:C / 1:3 形式")

    col_re = r"[A-Za-z]{1,3}"
    row_re = r"[1-9]\d{0,6}"
    single_re = re.compile(rf"^({col_re})({row_re})$")

    if ":" in raw:
        left, right = raw.split(":", 1)
        left, right = left.strip(), right.strip()
        # 整列 A:C
        if re.fullmatch(col_re, left) and re.fullmatch(col_re, right):
            c1, c2 = col_to_idx(left), col_to_idx(right)
            r1, r2 = 1, (ws.max_row if ws is not None else None)
            if r2 is None:
                raise CliError("bad_range", f"整列范围 {text} 需要提供工作表后才能解析")
            if r2 < 1:
                r2 = 1
            return _checked(r1, c1, r2, c2)
        # 整行 1:3
        if re.fullmatch(row_re, left) and re.fullmatch(row_re, right):
            r1, r2 = int(left), int(right)
            c1, c2 = 1, (ws.max_column if ws is not None else None)
            if c2 is None:
                raise CliError("bad_range", f"整行范围 {text} 需要提供工作表后才能解析")
            if c2 < 1:
                c2 = 1
            return _checked(r1, c1, r2, c2)
        m1, m2 = single_re.fullmatch(left), single_re.fullmatch(right)
        if not m1 or not m2:
            raise CliError("bad_range", f"无法解析范围 '{text}',应为 B2 / A1:B3 / A:C / 1:3 形式")
        r1, c1 = int(m1.group(2)), col_to_idx(m1.group(1))
        r2, c2 = int(m2.group(2)), col_to_idx(m2.group(1))
        return _checked(r1, c1, r2, c2)

    m = single_re.fullmatch(raw)
    if not m:
        raise CliError("bad_range", f"无法解析范围 '{text}',应为 B2 / A1:B3 / A:C / 1:3 形式")
    r, c = int(m.group(2)), col_to_idx(m.group(1))
    return _checked(r, c, r, c)


def _checked(r1: int, c1: int, r2: int, c2: int) -> tuple[int, int, int, int]:
    lo_r, hi_r = sorted((r1, r2))
    lo_c, hi_c = sorted((c1, c2))
    if hi_r > MAX_ROWS or hi_c > MAX_COLS or lo_r < 1 or lo_c < 1:
        raise CliError("bad_range", f"范围 {idx_to_col(c1)}{r1}:{idx_to_col(c2)}{r2} 超出 Excel 边界"
                                    f"(行 1-{MAX_ROWS},列 A-XFD)")
    return lo_r, lo_c, hi_r, hi_c


def area_label(r1: int, c1: int, r2: int, c2: int) -> str:
    a, b = f"{idx_to_col(c1)}{r1}", f"{idx_to_col(c2)}{r2}"
    return a if a == b else f"{a}:{b}"


def area_size(r1: int, c1: int, r2: int, c2: int) -> int:
    return (r2 - r1 + 1) * (c2 - c1 + 1)


# ---------------------------------------------------------------------------
# 值序列化:datetime -> ISO 字符串,保证 json.dumps 不出错
# ---------------------------------------------------------------------------

def serialize_value(v):
    """把 openpyxl 读到的值转为 JSON 安全的类型。

    - datetime/date/time -> ISO 字符串(如 2024-03-05T10:00:00 / 2024-03-05)
    - bool/number/str/None 原样;公式与错误值的缓存也是 str
    """
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, datetime):
        return v.isoformat(timespec="seconds")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.isoformat(timespec="seconds")
    if isinstance(v, (int, float, str)):
        return v
    return str(v)


def cell_type(cell: Cell) -> str:
    """单元格数据类别:formula / error / bool / number / date / str / blank"""
    dt = cell.data_type
    if dt == "f":
        return "formula"
    if dt == "e":
        return "error"
    if dt == "b":
        return "bool"
    if dt == "n":
        return "number"
    if dt == "d":
        return "date"
    if dt in ("s", "str", "inlineStr"):
        return "str"
    if cell.value is None:
        return "blank"
    return "str"


def is_formula_text(v) -> bool:
    """文本是否会被 openpyxl 当作公式(以 = 开头)"""
    return isinstance(v, str) and v.startswith("=")


# ---------------------------------------------------------------------------
# 工作簿加载 / 原子保存
# ---------------------------------------------------------------------------

def open_workbook(path: str, *, data_only: bool = False, read_only: bool = False) -> Workbook:
    """加载工作簿;统一错误信息。

    - .xls 旧格式直接给出可操作提示
    - .xlsm/.xltm 自动 keep_vba,避免保存时丢宏
    """
    if not os.path.exists(path):
        raise CliError("no_file", f"文件不存在: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        raise CliError(
            "unsupported_format",
            f"'{path}' 是旧版 .xls 格式,openpyxl 不支持。"
            f"请先用 Excel/WPS 另存为 .xlsx,或改用 xlrd 方案。",
        )
    if ext not in SUPPORTED_EXTS:
        raise CliError(
            "unsupported_format",
            f"不支持的文件类型 '{ext}',仅支持 {', '.join(sorted(SUPPORTED_EXTS))}",
        )
    try:
        keep_vba = ext in (".xlsm", ".xltm")
        return load_workbook(path, data_only=data_only, read_only=read_only, keep_vba=keep_vba)
    except CliError:
        raise
    except Exception as e:  # BadZipFile / KeyError / XML 解析错误等
        raise CliError("cannot_open", f"无法打开 {path}: {type(e).__name__}: {e}"
                                      f"(文件可能损坏,或不是有效的 Excel 文件)") from e


def save_workbook_atomic(wb: Workbook, path: str) -> None:
    """先写同目录临时文件再原子替换;避免中途失败损坏原文件。

    需要保留文件扩展名,因此临时文件使用同名不同前缀。
    Windows 上杀软/同步盘会对新落盘文件瞬时加锁,故对 PermissionError 退避重试;
    仍失败则最后直接写原文件一次(此时多半是 Excel 占用,直接写同样会失败)。
    """
    last_err = None
    waits = (0.3, 0.5, 0.8, 1.3, 2.0, 3.0)
    for attempt in range(len(waits) + 1):
        direct = attempt == len(waits)
        try:
            if direct:
                wb.save(path)
            else:
                _save_once(wb, path)
            return
        except PermissionError as e:
            last_err = e
            if not direct:
                _sleep(waits[attempt])

        except OSError as e:
            raise CliError("save_failed", f"写入 {path} 失败: {e}") from e
    raise CliError(
        "file_busy",
        f"无法写入 {path}: 文件可能正被 Excel/WPS/杀毒软件占用(已重试多次)。"
        f"请关闭占用程序后重试。",
    ) from last_err


def _save_once(wb: Workbook, path: str) -> None:
    directory = os.path.dirname(os.path.abspath(path))
    name = os.path.basename(path)
    fd, tmp_path = tempfile.mkstemp(prefix=".office-tmp-", suffix=name, dir=directory)
    os.close(fd)
    try:
        wb.save(tmp_path)
        replace_with_retry(tmp_path, path)
    except OSError:
        raise
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass


def choose_sheet(wb: Workbook, name: Optional[str]) -> Worksheet:
    """取工作表;name 缺省用打开时激活的表(与 Excel 行为一致)。"""
    if name is None:
        return wb.active
    try:
        return wb[name]
    except KeyError:
        raise CliError(
            "no_sheet",
            f"工作表 '{name}' 不存在。可用工作表: {wb.sheetnames}",
        ) from None


def data_dimensions(ws: Worksheet) -> Optional[tuple[int, int, int, int]]:
    """表内有数据的范围;空表(无任何值)返回 None。

    openpyxl 空表 max_row 可能为 1,因此按 A1 是否有值判断。
    """
    mr, mc = ws.max_row or 1, ws.max_column or 1
    if mr <= 1 and mc <= 1 and ws["A1"].value is None:
        return None
    return (1, 1, mr, mc)


def iter_area(ws: Worksheet, r1: int, c1: int, r2: int, c2: int):
    """按行遍历区域中的 Cell 对象(对不存在的格会新建空 Cell,不写盘即无副作用)。"""
    return ws.iter_rows(min_row=r1, min_col=c1, max_row=r2, max_col=c2)


def merged_anchor(ws, row: int, col: int) -> str | None:
    """(row, col) 落在某合并区且非左上角时返回左上角坐标(如 'A1'),否则 None。

    可用于写入前校验:合并区非左上格在 openpyxl 中是不可写的 MergedCell。
    """
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            if row == rng.min_row and col == rng.min_col:
                return None
            return ws.cell(row=rng.min_row, column=rng.min_col).coordinate
    return None


def sheet_names_snapshot(wb: Workbook) -> list[str]:
    return list(wb.sheetnames)


def guard_rowcol_shift(ws, axis: str, at: int, verb: str = "操作") -> None:
    """openpyxl 插入/删除行列不维护合并区/表格/筛选,凡会被移动/破坏的一律拒绝。

    axis: 'rows' | 'cols';at: 1-based 起点(删除带的起始行/列)。受影响判定:
    区域在起点及以下(右)即会错位或被破坏。
    """
    from .errors import CliError

    import re

    def ref_box(ref: str) -> tuple[int, int, int, int]:
        left, _, right = ref.partition(":")
        right = right or left
        m1 = re.match(r"([A-Z]+)(\d+)", left)
        m2 = re.match(r"([A-Z]+)(\d+)", right)
        if not m1 or not m2:
            raise CliError("internal", f"无法解析区域 {ref}")
        return (int(m1.group(2)), int(m2.group(2)),
                col_to_idx(m1.group(1)), col_to_idx(m2.group(1)))

    if axis == "rows":
        hits = [(str(rng), "合并单元格") for rng in ws.merged_cells.ranges
                if rng.max_row >= at]
        hits += [(t.name, "表格对象") for t in ws.tables.values()
                 if ref_box(t.ref)[1] >= at]
        if ws.auto_filter.ref is not None and ws.auto_filter.ref.max_row >= at:
            hits.append((str(ws.auto_filter.ref), "自动筛选区域"))
    else:
        hits = [(str(rng), "合并单元格") for rng in ws.merged_cells.ranges
                if rng.max_col >= at]
        hits += [(t.name, "表格对象") for t in ws.tables.values()
                 if ref_box(t.ref)[3] >= at]
        if ws.auto_filter.ref is not None and ws.auto_filter.ref.max_col >= at:
            hits.append((str(ws.auto_filter.ref), "自动筛选区域"))

    if hits:
        kinds = sorted({k for _, k in hits})
        raise CliError(
            "layout_conflict",
            f"{verb}会使已有{kinds}错位(openpyxl 不会自动移动它们);"
            f"请先 excel merge --unmerge 取消合并 / excel layout --unfilter 取消"
            f"筛选后重试")
