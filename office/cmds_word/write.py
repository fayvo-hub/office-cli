"""word write — 创建/追加 Word 文档(.docx),支持段落/标题/表格/图片/代码块。

- 文件不存在:需 --create(自动新建 .docx)
- 文件已存在:在末尾追加(rich 排版建议用 office md to-docx,此处适合程序化写入)
- 旧版 .doc:自动升级后追加,结果保存为同目录 .docx(原文件不动)
"""

from __future__ import annotations

import argparse
import json
import os

from .. import ioplan
from ..cli import add_file_arg
from ..errors import CliError

NAME = "write"
HELP = "创建/追加 Word 文档:段落、标题、表格、图片、代码块"
DESCRIPTION = """创建或追加 .docx 文档,数据用 JSON(建议 --data-file 传入,避免转义问题)。

--data-file 结构(顶层 dict):
{
  "blocks": [
    {"type": "h",    "level": 1, "text": "大标题"},            // 标题 1-9 级
    {"type": "p",    "text": "普通段落", "bold": true,
                     "size": 12, "color": "FF0000", "align": "center"},   // 可选样式
    {"type": "p",    "runs": [                                   // 富文本段落
                     {"text": "粗体", "bold": true},
                     {"text": "斜体", "italic": true}]},
    {"type": "code", "text": "print('hi')"},                    // 等宽字体代码块
    {"type": "quote", "text": "引用文字"},                       // 斜体缩进
    {"type": "table", "rows": [["列A", "列B"], ["1", "2"]], "header": true},
    {"type": "img",  "path": "C:/pic.png", "width": 12},        // width 单位 cm
    {"type": "pagebreak"}
  ]
}

说明:
- --text 快捷方式 = 追加一段普通文本(不能与 --data-file 同用)
- 图片路径相对当前工作目录解析;支持 png/jpg/gif
- 输出 JSON: {"ok": true, "file": "...", "appended_blocks": 8}
"""


def register(sp: argparse.ArgumentParser) -> None:
    add_file_arg(sp, help_text="Word 文件路径(.docx;旧版 .doc 存在时自动升级追加,输出 .docx)")
    src = sp.add_mutually_exclusive_group()
    src.add_argument("--data-file", metavar="PATH",
                     help="UTF-8 JSON 文件(结构见上;AI 首选)")
    src.add_argument("--text", metavar="TEXT", help="快捷方式:追加一段普通文本")
    sp.add_argument("--create", action="store_true",
                    help="文件不存在时新建;文件已存在时忽略本参数(直接追加)")
    sp.add_argument("--no-header", action="store_true", help="表格不把首行加粗(默认加粗)")


def run(args: argparse.Namespace) -> dict:
    if args.data_file is not None and not os.path.exists(args.data_file):
        raise CliError("no_file", f"--data-file 不存在: {args.data_file}")
    if args.text is None and args.data_file is None:
        raise CliError("bad_args", "请提供 --text 或 --data-file")

    try:
        from docx import Document
    except ImportError:  # pragma: no cover
        raise CliError("need_dep", "word 命令需要 python-docx: pip install python-docx") from None

    warnings: list[str] = []
    exists = os.path.exists(args.file)
    if args.file.lower().endswith(".doc"):
        if not exists:
            raise CliError("unsupported_format",
                           "不能创建旧版 .doc 文件,请把 --file 改为 .docx 后缀")
        plan = ioplan.word_plan(args.file)
        doc = Document(plan.read_path)
        warnings.append(f"{args.file} 为旧版 .doc,追加结果保存为 {plan.write_path}"
                        f"(原文件未改动)")
    elif exists:
        try:
            doc = Document(args.file)
        except Exception as e:
            raise CliError("cannot_open", f"无法打开 {args.file}: {e}") from e
        plan = None
    else:
        if not args.create:
            raise CliError("no_file", f"文件不存在: {args.file}"
                                      f"(新建 Word 文件请加 --create)")
        from docx import Document as _D

        doc = _D()
        plan = None
        warnings.append(f"已新建文件 {args.file}")

    blocks = _load_blocks(args)
    for b in blocks:
        _add_block(doc, b, header=not args.no_header)

    out_path = plan.write_path if plan else args.file
    try:
        doc.save(out_path)
    except PermissionError:
        raise CliError("file_busy", f"无法写入 {out_path}: 文件可能正被 WPS/Word 打开") from None
    except OSError as e:
        raise CliError("write_failed", f"写入 {out_path} 失败: {e}") from e

    _up = {"upgraded_from": plan.upgraded_from} if plan and plan.upgraded_from else {}
    return {"ok": True, "file": out_path, **_up,
            "appended_blocks": len(blocks), "warnings": warnings}


# ---------------------------------------------------------------------------
# 内部
# ---------------------------------------------------------------------------

def _load_blocks(args) -> list[dict]:
    if args.text is not None:
        return [{"type": "p", "text": args.text}]
    try:
        with open(args.data_file, encoding="utf-8") as fh:
            doc = json.load(fh)
    except (OSError, json.JSONDecodeError) as e:
        raise CliError("bad_json", f"读取 {args.data_file} 失败: {e}") from e
    if isinstance(doc, str):
        return [{"type": "p", "text": doc}]
    if isinstance(doc, list) and all(isinstance(b, dict) for b in doc):
        return doc
    if isinstance(doc, dict) and isinstance(doc.get("blocks"), list):
        return doc["blocks"]
    raise CliError("bad_json",
                   "结构不识别:顶层应是 {\"blocks\": [...]}、blocks 数组或字符串(见 --help)")


def _add_block(doc, b: dict, *, header: bool) -> None:
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    btype = b.get("type", "p")

    if btype == "h":
        doc.add_heading(str(b.get("text", "")), level=int(b.get("level", 1)))
        return
    if btype == "pagebreak":
        from docx.enum.text import WD_BREAK

        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        return
    if btype == "table":
        rows = b.get("rows") or []
        if not rows or not isinstance(rows[0], list):
            raise CliError("bad_json", "table 块需要 rows: 二维数组")
        tb = doc.add_table(rows=len(rows), cols=len(rows[0]))
        try:
            tb.style = "Table Grid"
        except Exception:  # pragma: no cover
            pass
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = tb.cell(i, j)
                cell.text = str(v)
                if header and i == 0:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.bold = True
        return
    if btype == "img":
        path = b.get("path", "")
        if not os.path.exists(path):
            raise CliError("no_file", f"img 块的图片不存在: {path}")
        p = doc.add_paragraph()
        run = p.add_run()
        try:
            run.add_picture(path, width=Cm(float(b.get("width", 15))))
        except Exception as e:
            raise CliError("bad_image", f"插入图片 {path} 失败: {e}") from e
        return

    # 段落类:h 之外统一走 add_paragraph,再套样式
    text = b.get("text") or b.get("runs", [])
    p = doc.add_paragraph()
    _style_para(p, b, Pt, RGBColor, WD_ALIGN_PARAGRAPH)
    if isinstance(text, str):
        _add_rich(p, [{"text": text}], b, Pt, RGBColor)
    else:
        _add_rich(p, text, b, Pt, RGBColor)
    if btype == "code":
        _shade_para(p)
    elif btype == "quote":
        p.paragraph_format.left_indent = Cm(1.0)


def _add_rich(p, runs: list, b: dict, Pt, RGBColor) -> None:
    """把 runs(或单文本)写入段落,支持 bold/italic/code/color/size。"""
    if not runs:
        return
    base_color = b.get("color")
    base_size = b.get("size")
    for r in runs:
        if isinstance(r, str):
            r = {"text": r}
        run = p.add_run(str(r.get("text", "")))
        run.bold = bool(r.get("bold", b.get("bold", False)))
        run.italic = bool(r.get("italic", b.get("italic", False)))
        if b.get("type") == "quote":
            run.italic = True
        color = r.get("color") or base_color
        size = r.get("size") or base_size
        if color:
            run.font.color.rgb = RGBColor.from_string(str(color))
        if size:
            run.font.size = Pt(float(size))
        if r.get("code") or b.get("type") == "code":
            _mono(run)


def _mono(run) -> None:
    from docx.oxml.ns import qn

    run.font.name = "Consolas"
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(attr), "Consolas")
    if run.font.size is None:
        from docx.shared import Pt

        run.font.size = Pt(9)


def _style_para(p, b: dict, Pt, RGBColor, WD_ALIGN_PARAGRAPH) -> None:
    align = b.get("align")
    if align:
        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                       "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "right": WD_ALIGN_PARAGRAPH.RIGHT,
                       "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(align)
    if b.get("type") == "quote":
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.space_after = Pt(6)


def _shade_para(p) -> None:
    """给段落加浅灰底纹。"""
    from docx.oxml.ns import nsdecls, qn
    from docx.oxml import parse_xml

    pPr = p._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F2F2F2"/>')
    pPr.append(shd)
