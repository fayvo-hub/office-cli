"""word replace — 在 .docx 中查找替换文本(模板占位符填充)。"""

from __future__ import annotations

import argparse
import json
import os
import re

from .. import ioplan
from ..cli import add_file_arg
from ..errors import CliError

NAME = "replace"
HELP = "替换文本/填充模板占位符(--data-file 键值对或 --find/--replace)"
DESCRIPTION = """在 .docx 全文(正文、表格、页眉页脚)中替换文本,常用于模板占位符填充。

用法示例(二选一):
  office word replace -f 合同.docx --data-file repl.json
      repl.json: {"${客户名}": "张三", "{日期}": "2026-03-01", "TODO": ""}
      键 = 要替换的原文(含占位符形态),值 = 替换文本
  office word replace -f 合同.docx --find '{{金额}}' --replace '12,800 元'
      (单条替换;shell 里花括号记得加引号)

说明:
- 段落内文本若被 Word 拆到多个 run,占位符本身完整时会先按 run 替换;
  跨 run 的占位符会合并段落文本再替换(此时该段样式统一为首 run 样式,
  输出中会以 rebuilt_paragraphs 统计并提示)
- 默认区分大小写(占位符替换场景);加 --ignore-case 不区分
- 键按长度从长到短替换,可放心同时给 "${name}" 与 "name"
- 旧版 .doc 自动升级为 .docx 后操作,原文件不动
输出 JSON: {"ok": true, "file": "...", "occurrences": N,
            "rebuilt_paragraphs": N}
"""


def register(sp: argparse.ArgumentParser) -> None:
    add_file_arg(sp)
    sp.add_argument("--sheet", help=argparse.SUPPRESS)  # word 无 sheet,仅防误用
    g = sp.add_mutually_exclusive_group(required=True)
    g.add_argument("--find", metavar="TEXT", help="查找文本(--replace 配合)")
    g.add_argument("--data-file", metavar="JSON",
                   help="JSON 文件: {原文: 替换文本, ...}")
    sp.add_argument("--replace", default="", metavar="TEXT",
                    help="替换为(--find 模式,默认空串=删除)")
    sp.add_argument("--ignore-case", action="store_true", help="不区分大小写")


def run(args: argparse.Namespace) -> dict:
    if args.data_file is not None:
        if not os.path.exists(args.data_file):
            raise CliError("no_file", f"--data-file 不存在: {args.data_file}")
        try:
            with open(args.data_file, "r", encoding="utf-8") as fh:
                raw = json.load(fh)
        except (OSError, json.JSONDecodeError) as e:
            raise CliError("bad_json", f"读取 {args.data_file} 失败: {e}") from None
        if not isinstance(raw, dict) or not all(
                isinstance(k, str) for k in raw):
            raise CliError("bad_json",
                           "--data-file 顶层应为 JSON 对象,键为字符串")
        mapping = {k: "" if v is None else str(v) for k, v in raw.items()}
    else:
        mapping = {args.find: args.replace}
    mapping = {k: v for k, v in mapping.items() if k}
    if not mapping:
        raise CliError("bad_args", "没有可替换的键(键不能为空字符串)")

    plan = ioplan.word_plan(args.file)
    if not os.path.exists(plan.read_path):
        raise CliError("no_file", f"文件不存在: {args.file}")
    try:
        from docx import Document
    except ImportError:
        raise CliError("need_dep",
                       "word 命令需要 python-docx: pip install python-docx") from None
    try:
        doc = Document(plan.read_path)
    except Exception as e:
        raise CliError("cannot_open", f"无法打开 {args.file}: {e}") from e

    flags = re.IGNORECASE if args.ignore_case else 0
    # 键按长度降序,避免 "${name}" 与 "name" 竞争时先替换短的
    ordered = sorted(mapping.items(), key=lambda kv: len(kv[0]), reverse=True)

    def apply(text: str) -> tuple[str, int]:
        n = 0
        for k, v in ordered:
            if args.ignore_case:
                pat = re.compile(re.escape(k), flags)
                text, cnt = pat.subn(lambda _m: v, text)
            else:
                cnt = text.count(k)
                if cnt:
                    text = text.replace(k, v)
            n += cnt
        return text, n

    occurrences = rebuilt = 0

    def process(paragraphs) -> None:
        nonlocal occurrences, rebuilt
        for p in paragraphs:
            if not p.runs:
                continue
            # 1) 逐 run 替换(保留样式)
            done = 0
            for r in p.runs:
                t2, n = apply(r.text)
                if n:
                    r.text = t2
                    done += n
            # 2) 段落合并文本里仍有键残留 -> 跨 run 重建(样式统一为首 run)
            whole = "".join(r.text for r in p.runs)
            t2, n = apply(whole)
            if n:
                if t2 != whole:
                    p.runs[0].text = t2
                    for r in p.runs[1:]:
                        r.text = ""
                    rebuilt += 1
                occurrences += done + n
            else:
                occurrences += done

    process(doc.paragraphs)

    def walk_table(t) -> None:
        seen = set()
        for row in t.rows:
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue  # 水平合并单元格会重复出现
                seen.add(id(cell._tc))
                process(cell.paragraphs)
                for nt in cell.tables:
                    walk_table(nt)

    for t in doc.tables:
        walk_table(t)

    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header,
                     sec.first_page_footer, sec.even_page_header,
                     sec.even_page_footer):
            try:
                if part.is_linked_to_previous:
                    continue
                process(part.paragraphs)
                for t in part.tables:
                    walk_table(t)
            except Exception:
                pass  # 某些模板节不存在时按无内容处理

    out_path = plan.write_path
    try:
        doc.save(out_path)
    except PermissionError:
        raise CliError("file_busy", f"无法写入 {out_path}: 文件可能正被 WPS/Word 打开") from None
    except Exception as e:
        raise CliError("write_failed", f"写入 {out_path} 失败: {e}") from e

    warnings = []
    if rebuilt:
        warnings.append(
            f"{rebuilt} 个段落跨 run 合并重建,样式统一为段落首 run 样式"
            f"(通常发生在占位符被 Word 拆开时)")
    _up = {"upgraded_from": plan.upgraded_from} if plan.upgraded_from else {}
    return {"ok": True, "file": out_path, **_up,
            "occurrences": occurrences, "rebuilt_paragraphs": rebuilt,
            "warnings": warnings}
