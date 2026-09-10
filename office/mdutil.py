"""mdutil — Markdown 渲染核心。

提供:
  md_to_html_string(md_path, css_path=None) -> str  完整自包含单文件 HTML
  md_to_html(md_path, html_path, css_path=None)     写 HTML 文件
  md_to_pdf(md_path, pdf_path, css_path=None)       HTML -> Chrome 打印 PDF
  md_to_docx(md_path, docx_path)                    近似转换 Word 文档

特性: mermaid 图渲染(内联 mermaid.min.js)、代码高亮(pygments,内联 CSS)、
      GFM 表格、标题层级、A4 打印版式。
"""

from __future__ import annotations

import html as _html
import os
import re
import tempfile
from html.parser import HTMLParser

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ._atomic import replace_with_retry
from .errors import CliError
from . import printer

_MERMAID_RE = re.compile(r"```mermaid[ \t]*\n(.*?)(?:\n```|```)", re.S)
_MMD_PLACEHOLDER = "@@OFFICE_MERMAID_{}@@"

_ASSETS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       "assets", "mermaid.min.js")

PAGE_CSS = """
@page { size: A4; margin: 0; }
html, body { margin: 0; padding: 0; }
body {
  font-family: "Segoe UI", "Microsoft YaHei", "PingFang SC", sans-serif;
  font-size: 13px; line-height: 1.75; color: #1f2328;
}
article { max-width: 72ch; margin: 0 auto; padding: 36px 46px; }
h1, h2, h3, h4, h5, h6 { font-weight: 600; line-height: 1.4; margin: 1.3em 0 0.6em; }
h1 { font-size: 1.75em; border-bottom: 1px solid #e5e7eb; padding-bottom: .3em; }
h2 { font-size: 1.4em; border-bottom: 1px solid #f0f1f3; padding-bottom: .25em; }
h3 { font-size: 1.2em; } h4 { font-size: 1.05em; }
p { margin: .55em 0; }
code {
  font-family: Consolas, "Courier New", monospace; font-size: .92em;
  background: #f0f1f3; border-radius: 4px; padding: 1px 5px;
}
pre { margin: .8em 0; }
pre code {
  display: block; padding: 12px 16px; overflow-x: auto;
  border-radius: 6px; background: #f6f8fa; line-height: 1.55;
  border: 1px solid #e5e7eb;
}
/* pygments 高亮作用域 */
.codehilite { margin: .8em 0; }
.codehilite pre { margin: 0; }
table { border-collapse: collapse; margin: .9em auto; font-size: .95em; }
th, td { border: 1px solid #d0d7de; padding: 5px 12px; text-align: left; }
th { background: #f6f8fa; font-weight: 600; }
blockquote {
  margin: .8em 0; padding: 2px 16px; color: #57606a;
  border-left: 4px solid #d0d7de; background: #fafbfc;
}
blockquote p { margin: .4em 0; }
hr { border: none; border-top: 1px solid #e5e7eb; margin: 1.6em 0; }
ul, ol { padding-left: 2em; margin: .5em 0; }
li { margin: .18em 0; }
a { color: #0969da; text-decoration: none; }
.mermaid { text-align: center; margin: 1em 0; }
.mermaid svg { max-width: 100%; }
"""


def _codehilite_css() -> str:
    try:
        from pygments.formatters import HtmlFormatter
        return HtmlFormatter().get_style_defs(".codehilite")
    except Exception:  # pragma: no cover
        return ""


def _read_md(md_path: str) -> str:
    if not os.path.exists(md_path):
        raise CliError("no_file", f"文件不存在: {md_path}")
    try:
        with open(md_path, "r", encoding="utf-8") as fh:
            return fh.read()
    except UnicodeDecodeError:
        try:
            with open(md_path, "r", encoding="gb18030") as fh:
                return fh.read()
        except Exception:
            raise CliError("bad_encoding", f"无法读取 {md_path}(需 UTF-8/GBK)") from None


def _render_md_body(text: str) -> tuple[str, list[str]]:
    """Markdown -> HTML 片段;mermaid 代码块单独抽出返回。"""
    mermaid_blocks: list[str] = []
    def _sub(m: re.Match) -> str:
        mermaid_blocks.append(m.group(1).strip())
        return f"\n{_MMD_PLACEHOLDER.format(len(mermaid_blocks) - 1)}\n"
    text2 = _MERMAID_RE.sub(_sub, text)
    try:
        import markdown
        body = markdown.markdown(
            text2,
            extensions=["extra", "codehilite", "toc"],
            extension_configs={
                "codehilite": {"guess_lang": False},
                "toc": {"permalink": False},
            },
            output_format="html5",
        )
    except Exception as e:  # pragma: no cover
        raise CliError("md_render", f"Markdown 渲染失败: {e}") from e
    for i, code in enumerate(mermaid_blocks):
        esc = _html.escape(code, quote=True)
        # 占位符会被 markdown 包进 <p>, 这里闭合 p 让 div 提升为块级顶层
        body = body.replace(_MMD_PLACEHOLDER.format(i),
                            f'</p><div class="mermaid">{esc}</div><p>')
    return body, mermaid_blocks


def _title_of(md_path: str, body: str) -> str:
    m = re.search(r"<h1[^>]*>(.*?)</h1>", body, re.S)
    if m:
        return re.sub(r"<[^>]+>", "", m.group(1)).strip()
    return os.path.splitext(os.path.basename(md_path))[0]


def _mermaid_js() -> str:
    if os.path.exists(_ASSETS):
        with open(_ASSETS, "r", encoding="utf-8") as fh:
            return fh.read()
    return ""


def md_to_html_string(md_path: str, css_path: str | None = None) -> str:
    """返回完整自包含 HTML 字符串。"""
    text = _read_md(md_path)
    body, _blocks = _render_md_body(text)
    extra_css = ""
    if css_path:
        if not os.path.exists(css_path):
            raise CliError("no_file", f"样式文件不存在: {css_path}")
        with open(css_path, "r", encoding="utf-8") as fh:
            extra_css = fh.read()
    title = _html.escape(_title_of(md_path, body))
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
{_codehilite_css()}
{PAGE_CSS}
{extra_css}
</style>
</head>
<body>
<article>
{body}
</article>
<script>
{_mermaid_js()}
</script>
<script>
window.addEventListener('DOMContentLoaded', function () {{
  if (window.mermaid && document.querySelectorAll('.mermaid').length) {{
    try {{
      mermaid.initialize({{ startOnLoad: true, securityLevel: 'loose' }});
    }} catch (e) {{ console.warn('mermaid init failed', e); }}
  }}
}});
</script>
</body>
</html>
"""


def md_to_html(md_path: str, html_path: str, css_path: str | None = None) -> None:
    out_dir = os.path.dirname(os.path.abspath(html_path))
    os.makedirs(out_dir, exist_ok=True)
    _write_text_atomic(html_path, md_to_html_string(md_path, css_path))


def _browser_path() -> str | None:
    for cand in (os.environ.get("OFFICE_CHROME"), os.environ.get("CHROME_EXECUTABLE"),
                 r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                 r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
                 r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
                 r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"):
        if cand and os.path.exists(cand):
            return cand
    return None


def _wait_mermaid(page, timeout_ms: int = 8000) -> bool:
    try:
        page.wait_for_function(
            """() => {
              const els = document.querySelectorAll('.mermaid');
              if (!els.length) return true;
              return document.querySelectorAll('.mermaid svg').length === els.length;
            }""", timeout=timeout_ms)
        return True
    except Exception:
        return False


def render_html_to_pdf(html_path: str, pdf_path: str) -> dict:
    """用系统 Chrome/Edge(playwright)把 HTML 打印为 A4 PDF。"""
    browser = _browser_path()
    if not browser:
        raise CliError("need_chrome",
                       "未找到 Chrome/Edge。请安装 Chrome,或设置环境变量 OFFICE_CHROME")
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:  # pragma: no cover
        raise CliError("need_dep",
                       "缺 playwright 库,无法渲染 PDF。完整版安装见"
                       " https://github.com/fayvo-hub/office-cli#安装"
                       "(pip 装 office-cli[mdpdf] 即可)或改用 office md to-docx/html") from None
    mermaid_ok = True
    try:
        # Chrome 启动同样会初始化打印子系统并读默认打印机: 期间隔离到本地虚拟打印机
        with printer.isolate(), sync_playwright() as p:
            browser_obj = p.chromium.launch(executable_path=browser)
            try:
                page = browser_obj.new_page()
                page.goto("file:///" + html_path.replace("\\", "/"),
                          wait_until="load")
                mermaid_ok = _wait_mermaid(page)
                page.pdf(
                    path=pdf_path, format="A4", print_background=True,
                    margin={"top": "1.55cm", "bottom": "1.4cm",
                            "left": "1.5cm", "right": "1.5cm"},
                    display_header_footer=True,
                    header_template="<span></span>",
                    footer_template=(
                        '<div style="font-size:8px;color:#8a8f98;width:100%;'
                        'text-align:center;">第 <span class="pageNumber"></span>'
                        ' 页 / 共 <span class="totalPages"></span> 页</div>'),
                )
            finally:
                browser_obj.close()
    except CliError:
        raise
    except Exception as e:  # pragma: no cover
        raise CliError("render_failed", f"PDF 渲染失败: {e}") from e
    return {"mermaid_ok": mermaid_ok}


def md_to_pdf(md_path: str, pdf_path: str, css_path: str | None = None) -> dict:
    """md -> HTML 临时文件 -> Chrome 打印 PDF。"""
    os.makedirs(os.path.dirname(os.path.abspath(pdf_path)), exist_ok=True)
    html_str = md_to_html_string(md_path, css_path)
    fd, tmp = tempfile.mkstemp(suffix=".html", prefix=".office-md-")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            fh.write(html_str)
        return render_html_to_pdf(tmp, pdf_path)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass


# ---------------------------------------------------------------------------
# md -> docx(近似转换:标题/段落/列表/表格/代码/图片/引用)
# ---------------------------------------------------------------------------

class _E:
    __slots__ = ("tag", "attrs", "children")

    def __init__(self, tag: str, attrs: dict, children: list):
        self.tag = tag
        self.attrs = attrs
        self.children = children


_VOID = {"br", "hr", "img", "meta", "link", "input", "col", "wbr"}


class _TreeBuilder(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root: list = []
        self.stack: list[_E] = []
        self.cur: list = self.root  # 当前容器(children list)

    def handle_starttag(self, tag, attrs):
        node = _E(tag, dict(attrs), [])
        self.cur.append(node)
        if tag not in _VOID:
            self.stack.append(node)
            self.cur = node.children

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in _VOID and self.stack:
            self.stack.pop()
            self.cur = self.stack[-1].children if self.stack else self.root

    def handle_endtag(self, tag):
        if tag in _VOID:
            return
        if self.stack:
            self.stack.pop()
            self.cur = self.stack[-1].children if self.stack else self.root

    def handle_data(self, data):
        self.cur.append(data)


def _set_east_asia(style_or_run, name: str = "等线") -> None:
    """让中文使用指定字体(否则 Word 显示默认主题字体)。"""
    try:
        rpr = style_or_run.font.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), name)
    except Exception:
        pass


def _inline_runs(node, out: list[tuple[str, bool, bool, bool]]) -> None:
    """收集段内文本: (text, bold, italic, code)。图片单独走 _find_imgs。"""
    if isinstance(node, str):
        out.append((node, False, False, False))
        return
    if node.tag == "br":
        out.append(("\n", False, False, False))
        return
    if node.tag == "img":
        return  # 图片由 _collect 收集, 段落结束后另起图段
    bold = node.tag in ("strong", "b")
    italic = node.tag in ("em", "i")
    code = node.tag == "code"
    for ch in node.children:
        if isinstance(ch, str):
            out.append((ch, bold, italic, code))
        else:
            _inline_runs(ch, out)


def _find_imgs(node, img_dir: str, warnings: list[str]) -> list[tuple[str, int | None]]:
    """深搜 img 节点, 返回 [(绝对路径, width_px)]。"""
    found: list[tuple[str, int | None]] = []
    if isinstance(node, str):
        return found
    if node.tag == "img":
        src = node.attrs.get("src", "")
        if src:
            path = src if os.path.isabs(src) else os.path.join(img_dir, src)
            if not os.path.exists(path):
                warnings.append(f"图片不存在, 已跳过: {src}")
                return found
            try:
                width = int(float(node.attrs.get("width", 0)))
            except ValueError:
                width = 0
            found.append((path, width or None))
        return found
    for ch in node.children:
        found.extend(_find_imgs(ch, img_dir, warnings))
    return found


def _add_images(doc, imgs: list[tuple[str, int | None]]) -> None:
    for path, width in imgs:
        try:
            if width:
                pic = doc.add_picture(path, width=Cm(min(width / 96 * 2.54, 15.0)))
            else:
                pic = doc.add_picture(path, width=Cm(12.0))
            doc.paragraphs[-1].alignment = 1  # CENTER
        except Exception as e:
            raise CliError("img_fail", f"插入图片 {path} 失败: {e}") from e


def _add_code_lines(doc, text: str) -> None:
    """多行代码 -> 等宽灰底段落(每行一个 run + break)。"""
    p = doc.add_paragraph()
    lines = text.rstrip("\n").split("\n")
    for i, ln in enumerate(lines):
        run = p.add_run(ln)
        run.font.name = "Consolas"
        _set_east_asia(run, "Consolas")
        run.font.size = Pt(9.5)
        if i < len(lines) - 1:
            p.add_run().add_break()
    _shade(p, "F2F2F2")


def _add_para(doc, runs: list[tuple[str, bool, bool, bool]]) -> None:
    p = doc.add_paragraph()
    for text, bold, italic, code in runs:
        if not text:
            continue
        parts = text.split("\n")
        for i, part in enumerate(parts):
            if i:
                p.add_run().add_break()
            run = p.add_run(part)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if code:
                run.font.name = "Consolas"
                _set_east_asia(run, "Consolas")
                run.font.size = Pt(9.5)
    return p


_BLOCK_TAGS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote",
              "ul", "ol", "li", "table", "thead", "tbody", "tr", "td", "th",
              "div", "article"}


def _strip_blank(nodes: list) -> None:
    """删除块级容器间纯空白文本节点(如 HTML 换行), 防污染 docx 段落;
   同时移除因此产生的空 <p>。"""
    for node in nodes:
        if isinstance(node, _E) and node.tag in _BLOCK_TAGS:
            node.children = [c for c in node.children
                             if not (isinstance(c, str) and not c.strip())]
            _strip_blank(node.children)
    # 顶层/嵌套后: 去掉没有任何内容的空 p(mermaid div 提升留下的 <p></p>)
    i = 0
    while i < len(nodes):
        n = nodes[i]
        if isinstance(n, _E) and n.tag == "p" and not n.children:
            del nodes[i]
        else:
            i += 1


def _find_tags(node, want: str) -> list:
    """深搜指定标签(命中即不再下钻, 防止嵌套表格串行)。"""
    if isinstance(node, str):
        return []
    if node.tag == want:
        return [node]
    out: list = []
    for ch in node.children:
        if not isinstance(ch, str):
            out.extend(_find_tags(ch, want))
    return out


def _emit_docx(node, doc, img_dir: str, warnings: list[str]) -> None:
    if isinstance(node, str):
        return
    tag = node.tag
    if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
        runs, imgs = _collect(node, img_dir, warnings)
        _add_images(doc, imgs)
        text = "".join(t for t, *_ in runs)
        level = int(tag[1])
        if text.strip():
            try:
                h = doc.add_heading(level=min(level, 9))
                h.add_run(text)
                for r in h.runs:
                    _set_east_asia(r)
            except Exception:
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True
        return
    if tag == "p":
        runs, imgs = _collect(node, img_dir, warnings)
        _add_para(doc, runs)
        _add_images(doc, imgs)
        return
    if tag == "blockquote":
        runs, imgs = _collect(node, img_dir, warnings)
        _add_images(doc, imgs)
        p = _add_para(doc, runs)
        for r in p.runs:
            r.italic = True
            r.font.color.rgb = RGBColor(0x57, 0x60, 0x6A)
        p.paragraph_format.left_indent = Cm(0.6)
        return
    if tag in ("pre", "code"):
        _add_code_lines(doc, _text_of(node))
        return
    if tag == "div":
        if node.attrs.get("class", "").find("mermaid") >= 0:
            # mermaid 图在 docx 中输出源码代码块
            _add_code_lines(doc, _text_of(node))
            return
        for ch in node.children:
            _emit_docx(ch, doc, img_dir, warnings)
        return
    if tag in ("ul", "ol"):
        num = tag == "ol"
        for li in node.children:
            if isinstance(li, str):
                continue
            runs, imgs = _collect(li, img_dir, warnings)
            _add_images(doc, imgs)
            p = _add_para(doc, runs)
            p.style = doc.styles["List Number"] if num else doc.styles["List Bullet"]
        return
    if tag == "table":
        rows: list[list[str]] = []
        for tr in _find_tags(node, "tr"):
            cells = []
            for td in _find_tags(tr, "td") or _find_tags(tr, "th"):
                runs, _imgs = _collect(td, img_dir, warnings)
                cells.append(("".join(t for t, *_ in runs)).strip())
            rows.append(cells)
        if rows:
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    cell = t.cell(i, j)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    run = par.add_run(val)
                    if i == 0:
                        run.bold = True
                        _shade_cell(cell, "F2F2F2")
            doc.add_paragraph()
        return
    if tag == "hr":
        return
    if tag == "img":
        _add_images(doc, _find_imgs(node, img_dir, warnings))
        return
    if tag == "div":
        if node.attrs.get("class", "").find("mermaid") >= 0:
            # mermaid 图在 docx 中输出源码代码块
            _add_code_lines(doc, _text_of(node))
            return
        # 其它 div(如 codehilite 容器): 逐子递归
        for ch in node.children:
            _emit_docx(ch, doc, img_dir, warnings)
        return
    # 其它标签(span/a/strong 等块级外):逐子递归
    for ch in node.children:
        _emit_docx(ch, doc, img_dir, warnings)


def _collect(node, img_dir: str, warnings: list[str]):
    """收集节点内 (runs, imgs)。"""
    runs: list = []
    imgs = _find_imgs(node, img_dir, warnings)
    _inline_runs(node, runs)
    return runs, imgs


def _text_of(node) -> str:
    if isinstance(node, str):
        return node
    return "".join(_text_of(c) for c in node.children)


def _shade(p, hex_color: str) -> None:
    try:
        pPr = p._p.get_or_add_pPr()
        shd = pPr.makeelement(qn("w:shd"), {
            qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color})
        pPr.append(shd)
    except Exception:
        pass


def _shade_cell(cell, hex_color: str) -> None:
    try:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {
            qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color})
        tcPr.append(shd)
    except Exception:
        pass


def md_to_docx(md_path: str, docx_path: str) -> None:
    text = _read_md(md_path)
    body, _blocks = _render_md_body(text)
    builder = _TreeBuilder()
    builder.feed(body)
    _strip_blank(builder.root)
    os.makedirs(os.path.dirname(os.path.abspath(docx_path)), exist_ok=True)
    img_dir = os.path.dirname(os.path.abspath(md_path))
    warnings: list[str] = []
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    _set_east_asia(st)
    for node in builder.root:
        _emit_docx(node, doc, img_dir, warnings)
    _write_docx_atomic(doc, docx_path)


def _write_text_atomic(path: str, content: str) -> None:
    fd, tmp = tempfile.mkstemp(suffix=".tmp", prefix=".office-",
                               dir=os.path.dirname(os.path.abspath(path)))
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            fh.write(content)
        replace_with_retry(tmp, path)
    except Exception:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise


def _write_docx_atomic(doc: Document, path: str) -> None:
    fd, tmp = tempfile.mkstemp(suffix=".docx", prefix=".office-",
                               dir=os.path.dirname(os.path.abspath(path)))
    os.close(fd)
    try:
        doc.save(tmp)
        replace_with_retry(tmp, path)
    except Exception:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise
