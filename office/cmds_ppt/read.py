"""ppt read — 读取演示文稿(.pptx/.ppt)内容结构。"""

from __future__ import annotations

import argparse
import os

from .. import ioplan
from ..cli import add_file_arg
from ..errors import CliError

NAME = "read"
HELP = "读取 PPT:每页标题/文本/表格/图片/备注"
DESCRIPTION = """读取演示文稿(.pptx;旧版 .ppt 自动经 WPS 升级后读取),输出每页结构。

输出 JSON:
{
  "ok": true, "file": "...", "slides_total": 3, "pages": 3,
  "slides": [
    {"index": 1, "layout": "标题和内容", "title": "标题",
     "texts": ["要点一", {"text": "子要点", "level": 1}],
     "tables": [[["列1", "列2"], ["a", "1"]]],
     "pictures": 1, "charts": 0, "notes": "备注文本"},
    ...
  ],
  "warnings": []
}

说明:
- texts 为除标题外所有文本框内容(正文占位符与自由文本框),逐段输出;
  非顶层段落为 {"text": "...", "level": n}(0 = 顶层)
- tables 为该页所有表格(二维数组,单元格取全部文本)
- 图片/图表默认只给数量;加 --save-images DIR 可导出全部图片
- 加 --slide N 只输出第 N 页(1 起)
- 备注放 notes(无备注为 null)
"""


def register(sp: argparse.ArgumentParser) -> None:
    add_file_arg(sp, help_text="PPT 文件路径(.pptx;旧版 .ppt 自动升级)")
    sp.add_argument("--slide", type=int, metavar="N",
                    help="只输出第 N 页(默认全部)")
    sp.add_argument("--save-images", metavar="DIR",
                    help="把幻灯片内图片导出到目录(默认只统计数量)")


def run(args: argparse.Namespace) -> dict:
    plan = ioplan.ppt_plan(args.file, write=False)
    try:
        from pptx import Presentation
    except ImportError:  # pragma: no cover
        raise CliError("need_dep", "ppt 命令需要 python-pptx: pip install python-pptx") from None

    try:
        prs = Presentation(plan.read_path)
    except Exception as e:
        raise CliError("cannot_open", f"无法打开 {args.file}: {e}") from e

    warnings: list[str] = []
    if plan.upgraded_from:
        warnings.append(f"{plan.upgraded_from} 为旧版 .ppt,已自动升级读取(原文件未改动)")

    total = len(prs.slides)
    wanted = range(total)
    if args.slide is not None:
        if not 1 <= args.slide <= total:
            raise CliError("bad_args",
                           f"--slide 越界: {args.slide}(共 {total} 页)")
        wanted = range(args.slide - 1, args.slide)

    saved: list[str] = []
    slides_out = slides_of(prs, wanted)

    if args.save_images:
        _export_images(prs, args.save_images, saved)

    _up = {"upgraded_from": plan.upgraded_from} if plan.upgraded_from else {}
    return {"ok": True, "file": args.file, **_up,
            "slides_total": total, "pages": len(slides_out),
            "slides": slides_out,
            **({"images_saved": saved} if args.save_images else {}),
            "warnings": warnings}


def slides_of(prs, wanted=None) -> list[dict]:
    """演示文稿 → 每页结构列表(rag prep 与 ppt read 共用)。

    返回 [{index, layout, title, texts, tables, pictures, charts, notes}];
    wanted 传 range/可迭代时只取对应 0-based 页。
    """
    if wanted is None:
        wanted = range(len(prs.slides))
    slides_out = []
    for idx in wanted:
        slide = prs.slides[idx]
        title_ph = slide.shapes.title
        title = None
        if title_ph is not None and title_ph.has_text_frame:
            title = title_ph.text_frame.text.strip() or None
        texts: list = []
        tables = []
        pictures = charts = 0
        for shape in slide.shapes:
            if title_ph is not None and shape.shape_id == title_ph.shape_id:
                continue  # 标题占位符已单独放入 title
            try:
                st = shape.shape_type
            except Exception:
                st = None
            if st == 13:  # MSO_SHAPE_TYPE.PICTURE
                pictures += 1
                continue
            if getattr(shape, "has_chart", False):
                charts += 1
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = (para.text or "").strip()
                    if not txt:
                        continue
                    texts.append({"text": txt, "level": para.level}
                                 if para.level else txt)
            elif shape.has_table:
                rows = [[c.text for c in row.cells]
                        for row in shape.table.rows]
                tables.append(rows)
        notes = _notes(slide)
        slides_out.append({
            "index": idx + 1,
            "layout": _layout_name(slide.slide_layout),
            "title": title,
            "texts": texts,
            "tables": tables,
            "pictures": pictures,
            "charts": charts,
            "notes": notes,
        })
    return slides_out


# ---------------------------------------------------------------------------
# 内部
# ---------------------------------------------------------------------------

def _layout_name(layout) -> str:
    try:
        return layout.name or ""
    except Exception:
        return ""


def _notes(slide) -> str | None:
    try:
        if not slide.has_notes_slide:
            return None
        txt = slide.notes_slide.notes_text_frame.text.strip()
        return txt or None
    except Exception:
        return None


def _export_images(prs, out_dir: str, saved: list[str]) -> None:
    os.makedirs(out_dir, exist_ok=True)
    n = 0
    for si, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            try:
                if shape.shape_type != 13 or shape.image is None:
                    continue
            except Exception:
                continue
            n += 1
            try:
                img = shape.image
                name = f"slide{si}-{n}.{img.ext or 'png'}"
                with open(os.path.join(out_dir, name), "wb") as fh:
                    fh.write(img.blob)
                saved.append(name)
            except Exception:
                continue
