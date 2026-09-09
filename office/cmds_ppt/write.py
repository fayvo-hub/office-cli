"""ppt write — 新建/追加演示文稿(.pptx)。"""

from __future__ import annotations

import argparse
import json
import os

from .. import ioplan
from ..cli import add_file_arg
from ..errors import CliError

NAME = "write"
HELP = "创建/追加 PPT 幻灯片(标题/要点/表格/图片/备注)"
DESCRIPTION = """创建新演示文稿(加 --create)或向现有 .pptx 末尾追加幻灯片。

数据用 --data-file 提供 UTF-8 JSON(AI 首选),结构:
{
  "slides": [
    {
      "layout": "title|title-content|blank",     // 缺省自动推断
      "title": "页标题",
      "bullets": ["第一点", {"text": "子要点", "level": 1}],  // 与 body 二选一
      "body": "单段文本(放入正文占位符)",
      "notes": "演讲者备注",
      "texts": [{"text": "...", "size_pt": 24}], // 自由文本框(仅 blank 页)
      "table": {"data": [["列1","列2"],["a","1"]],
                "left_in": 0.6, "top_in": 1.8, "width_in": 8.8},
      "picture": {"path": "logo.png", "left_in": 0.6, "top_in": 2.0,
                  "width_in": 3.0}
    }
  ]
}

规则:
- layout 缺省自动:含 bullets/body -> title-content;仅 title -> title;
  含 table/picture/texts -> blank(此时 layout 参数被忽略并记录警告)
- 所有尺寸单位为英寸(默认页 10×7.5);位置缺省自动摆放,先给 title/table/
  picture 再给 texts(自上而下堆叠)
- 图片路径相对当前工作目录;建议 png/jpg
- 无 --create 时打开现有文件,幻灯片追加到末尾
输出 JSON: {"ok": true, "file": "...", "created": true/false,
            "slides_added": N, "warnings": []}
"""


def register(sp: argparse.ArgumentParser) -> None:
    add_file_arg(sp, help_text="PPT 文件路径(.pptx;旧版 .ppt 存在时自动升级追加,输出 .pptx)")
    sp.add_argument("--data-file", required=True, metavar="PATH",
                    help="UTF-8 JSON 文件(结构见上)")
    sp.add_argument("--create", action="store_true",
                    help="文件不存在时新建;已存在时忽略本参数(直接追加)")


def run(args: argparse.Namespace) -> dict:
    if not os.path.exists(args.data_file):
        raise CliError("no_file", f"--data-file 不存在: {args.data_file}")
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:  # pragma: no cover
        raise CliError("need_dep", "ppt 命令需要 python-pptx: pip install python-pptx") from None

    with open(args.data_file, encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, dict) or not isinstance(data.get("slides"), list) \
            or not data["slides"]:
        raise CliError("bad_args", "--data-file 需含非空 slides 数组")

    warnings: list[str] = []
    exists = os.path.exists(args.file)
    if args.file.lower().endswith(".ppt"):
        if not exists:
            raise CliError("unsupported_format",
                           "不能创建旧版 .ppt 文件,请把 --file 改为 .pptx 后缀")
        plan = ioplan.ppt_plan(args.file)
        prs = Presentation(plan.read_path)
        warnings.append(f"{args.file} 为旧版 .ppt,追加结果保存为 {plan.write_path}"
                        f"(原文件未改动)")
    elif exists:
        try:
            prs = Presentation(args.file)
        except Exception as e:
            raise CliError("cannot_open", f"无法打开 {args.file}: {e}") from e
        plan = None
    else:
        if not args.create:
            raise CliError("no_file", f"文件不存在: {args.file}"
                                      f"(新建 PPT 文件请加 --create)")
        prs = Presentation()
        plan = None
        warnings.append(f"已新建文件 {args.file}")

    in_ = Inches
    made = 0
    for sd in data["slides"]:
        if not isinstance(sd, dict):
            raise CliError("bad_args", f"slide 项需为对象: {sd!r}")
        _add_slide(prs, sd, warnings, in_, Pt)
        made += 1

    out_path = plan.write_path if plan else args.file
    try:
        prs.save(out_path)
    except PermissionError:
        raise CliError("file_busy", f"无法写入 {out_path}: 文件可能正被 WPS/WPS 打开") from None
    except OSError as e:
        raise CliError("write_failed", f"写入 {out_path} 失败: {e}") from e

    _up = {"upgraded_from": plan.upgraded_from} if plan and plan.upgraded_from else {}
    return {"ok": True, "file": out_path, **_up, "created": not exists,
            "slides_added": made, "warnings": warnings}


# ---------------------------------------------------------------------------
# 内部
# ---------------------------------------------------------------------------

_LAYOUT_ALIASES = {"title": 0, "title-content": 1, "blank": 6}


def _pick_layout(prs, sd: dict, warnings: list[str]) -> object:
    """决定版式:显式 layout 优先;含 table/picture/texts 强制 blank。"""
    has_free = any(k in sd for k in ("table", "picture", "texts"))
    explicit = sd.get("layout")
    if has_free and explicit not in (None, "blank"):
        warnings.append("slide 含 table/picture/texts,已改用 blank 版式手动排版"
                        "(原 layout 参数忽略)")
    if has_free:
        return prs.slide_layouts[6]
    if explicit is None:
        if sd.get("bullets") or sd.get("body"):
            return prs.slide_layouts[1]
        if sd.get("title"):
            return prs.slide_layouts[0]
        return prs.slide_layouts[6]
    if isinstance(explicit, int) or isinstance(explicit, str) \
            and str(explicit).isdigit():
        n = int(explicit)
        if not 0 <= n < len(prs.slide_layouts):
            raise CliError("bad_args", f"layout 索引越界: {n}"
                                       f"(模板共 {len(prs.slide_layouts)} 个版式)")
        return prs.slide_layouts[n]
    n = _LAYOUT_ALIASES.get(str(explicit).lower())
    if n is None:
        raise CliError("bad_args",
                       f"layout 未知: {explicit}(可用 title/title-content/blank 或版式索引)")
    return prs.slide_layouts[n]


def _add_slide(prs, sd: dict, warnings: list[str], in_, Pt) -> None:
    layout = _pick_layout(prs, sd, warnings)
    slide = prs.slides.add_slide(layout)
    title = sd.get("title")
    has_free = any(k in sd for k in ("table", "picture", "texts"))

    cursor = 0.45  # blank 页手动排版的当前纵坐标(英寸)
    if title:
        _put_title(slide, layout, title, has_free, in_, Pt)
        cursor = 1.65

    if sd.get("bullets") or sd.get("body"):
        _put_body(slide, sd, in_, Pt)
    if sd.get("notes"):
        try:
            slide.notes_slide.notes_text_frame.text = str(sd["notes"])
        except Exception as e:
            warnings.append(f"备注写入失败: {e}")

    cursor = _put_free(slide, sd, cursor, in_, Pt, warnings)
    if title and not has_free:
        cursor += 0.15


def _put_title(slide, layout, title: str, has_free: bool, in_, Pt) -> None:
    if not has_free and slide.shapes.title is not None:
        try:
            slide.shapes.title.text = title
            return
        except Exception:
            pass
    # blank 页或占位符不可用:手动顶部文本框
    box = slide.shapes.add_textbox(in_(0.45), in_(0.4), in_(9.1), in_(1.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(30)
    r.font.bold = True


def _put_body(slide, sd: dict, in_, Pt) -> None:
    ph = None
    for sh in slide.placeholders:
        try:
            if sh.placeholder_format.idx == 1:
                ph = sh
                break
        except Exception:
            continue
    items = sd.get("bullets")
    if ph is not None:
        tf = ph.text_frame
        if items:
            first = True
            for it in items:
                if isinstance(it, dict):
                    txt, lv = str(it.get("text", "")), int(it.get("level", 0))
                else:
                    txt, lv = str(it), 0
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.text = txt
                p.level = lv
        else:
            tf.text = str(sd.get("body", ""))
    else:
        # 无内容占位符(如 title 版式):直接加文本框兜底
        box = slide.shapes.add_textbox(in_(0.45), in_(1.7), in_(9.1), in_(5.0))
        tf = box.text_frame
        tf.word_wrap = True
        if items:
            first = True
            for it in items:
                txt, lv = (str(it.get("text", "")), int(it.get("level", 0))) \
                    if isinstance(it, dict) else (str(it), 0)
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.text = txt
                p.level = lv
        else:
            tf.text = str(sd.get("body", ""))


def _put_free(slide, sd: dict, cursor: float, in_, Pt,
              warnings: list[str]) -> float:
    """blank 页自由元素:table -> picture -> texts(自上而下)。返回最终纵坐标。"""
    W = 9.1  # 可用宽度(默认页宽 10 英寸留边)
    if "table" in sd:
        tbl = sd["table"]
        data = tbl.get("data")
        if not isinstance(data, list) or not data \
                or not isinstance(data[0], list):
            raise CliError("bad_args", "table.data 需为二维数组")
        rows, cols = len(data), max(len(r) for r in data)
        left = float(tbl.get("left_in", 0.45))
        top = float(tbl.get("top_in", cursor))
        width = float(tbl.get("width_in", W))
        if not 0.5 <= width <= 10:
            raise CliError("bad_args", f"table.width_in 需在 0.5..10,收到 {width}")
        gt = slide.shapes.add_table(rows, cols, in_(left), in_(top),
                                    in_(width), in_(0.4 * rows)).table
        if cols:
            col_w = int(in_(width) / cols)  # 等分总宽(EMU)
            for c in range(cols):
                gt.columns[c].width = col_w
        for r in range(rows):
            for c in range(cols):
                val = data[r][c] if c < len(data[r]) else ""
                gt.cell(r, c).text = "" if val is None else str(val)
        cursor = top + 0.4 * rows + 0.3
    if "picture" in sd:
        pic = sd["picture"]
        path = pic.get("path")
        if not path or not os.path.exists(path):
            raise CliError("no_file", f"picture.path 不存在: {path}")
        left = float(pic.get("left_in", 0.45))
        top = float(pic.get("top_in", cursor))
        width = float(pic.get("width_in", 4.0))
        if not 0.2 <= width <= 10:
            raise CliError("bad_args", f"picture.width_in 需在 0.2..10,收到 {width}")
        try:
            slide.shapes.add_picture(path, in_(left), in_(top), width=in_(width))
        except Exception as e:
            raise CliError("bad_args", f"图片插入失败: {e}") from None
        cursor = top + width * 0.75 + 0.3
    if "texts" in sd:
        for it in sd["texts"]:
            txt = str(it.get("text", ""))
            if not txt:
                continue
            left = float(it.get("left_in", 0.45))
            top = float(it.get("top_in", cursor))
            width = float(it.get("width_in", W))
            height = float(it.get("height_in", 0.6))
            box = slide.shapes.add_textbox(in_(left), in_(top),
                                           in_(width), in_(height))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(int(it.get("size_pt", 18)))
            if it.get("bold"):
                r.font.bold = True
            cursor = top + height + 0.15
    return cursor
