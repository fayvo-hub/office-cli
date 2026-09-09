"""layout — 工作表版式:列宽/行高/冻结窗格/自动筛选/隐藏行列。"""

from __future__ import annotations

import argparse
import re

from openpyxl.utils import get_column_letter

from .. import ioplan, xlutil
from ..cli import add_file_arg, add_sheet_arg
from ..errors import CliError

NAME = "layout"
HELP = "版式:列宽/行高/冻结窗格/自动筛选/隐藏行列"
DESCRIPTION = """工作表版式设置(不影响单元格内容)。

用法示例:
  office layout -f f.xlsx --col-width 'A=18,C:E=22' --row-height '1=28,3:5=20'
  office layout -f f.xlsx --freeze A2            # 冻结:滚动时首行(及以上/左侧)固定
  office layout -f f.xlsx --unfreeze
  office layout -f f.xlsx --filter A1:F100       # 表头行加自动筛选(下拉箭头)
  office layout -f f.xlsx --unfilter
  office layout -f f.xlsx --hide-cols 'C:F' --hide-rows '2:5'
  office layout -f f.xlsx --show-cols 'C:F' --show-rows '2:5'   # 取消隐藏

参数约定:
- --col-width / --row-height:逗号分隔的多段规格,每段 '范围=数值',
  列范围如 A、C:E;行范围如 1、3:5;数值单位:列宽=字符数,行高=磅
- --freeze REF:REF 指滚动时保持固定的边界格: A2=冻结第1行,
  B1=冻结A列, B2=同时冻结首行+A列;写 A1 等价于取消
- --filter 区域需包含表头行;被过滤的行隐藏由用户在 Excel 里操作,
  本命令只建立筛选状态
输出 JSON: {"ok": true, "file": "...", "applied": [...]}
"""


def register(sp: argparse.ArgumentParser) -> None:
    add_file_arg(sp)
    add_sheet_arg(sp)
    sp.add_argument("--col-width", metavar="SPEC",
                    help="列宽,如 'A=18,C:E=22'(字符数;多段逗号分隔)")
    sp.add_argument("--row-height", metavar="SPEC",
                    help="行高,如 '1=28,3:5=20'(磅;多段逗号分隔)")
    fz = sp.add_mutually_exclusive_group()
    fz.add_argument("--freeze", metavar="REF",
                    help="冻结窗格,如 A2(冻结首行)/B1(冻结A列)/B2")
    fz.add_argument("--unfreeze", action="store_true", help="取消冻结")
    fl = sp.add_mutually_exclusive_group()
    fl.add_argument("--filter", metavar="REF",
                    help="自动筛选区域(须含表头行),如 A1:F100")
    fl.add_argument("--unfilter", action="store_true", help="取消自动筛选")
    sp.add_argument("--hide-cols", metavar="RANGE", help="隐藏列,如 C:F")
    sp.add_argument("--show-cols", metavar="RANGE", help="取消隐藏列")
    sp.add_argument("--hide-rows", metavar="RANGE", help="隐藏行,如 2:5")
    sp.add_argument("--show-rows", metavar="RANGE", help="取消隐藏行")


def run(args: argparse.Namespace) -> dict:
    if not any(v is not None and v is not False
               for v in (args.col_width, args.row_height, args.freeze,
                         args.unfreeze, args.filter, args.unfilter,
                         args.hide_cols, args.show_cols,
                         args.hide_rows, args.show_rows)):
        raise CliError("bad_args", "未指定任何版式操作(--col-width/--row-height/"
                                   "--freeze/--filter/--hide-cols 等)")

    plan = ioplan.excel_plan(args.file)
    wb = xlutil.open_workbook(plan.read_path)
    ws = xlutil.choose_sheet(wb, args.sheet)
    applied: list[str] = []

    # ---- 列宽 / 行高 ----
    if args.col_width is not None:
        n = 0
        for c1, c2, val in _parse_specs(args.col_width, axis="col"):
            for ci in range(c1, c2 + 1):
                ws.column_dimensions[get_column_letter(ci)].width = val
                n += 1
        applied.append(f"col_width x{n}")

    if args.row_height is not None:
        n = 0
        for r1, r2, val in _parse_specs(args.row_height, axis="row"):
            for ri in range(r1, r2 + 1):
                ws.row_dimensions[ri].height = val
                n += 1
        applied.append(f"row_height x{n}")

    # ---- 冻结窗格 ----
    if args.unfreeze:
        ws.freeze_panes = None
        applied.append("unfreeze")
    elif args.freeze is not None:
        r1, c1, r2, c2 = xlutil.parse_range(args.freeze, ws)
        # 只接受单个单元格锚点(如 A2/B1/B2)
        if not (r1 == r2 and c1 == c2):
            raise CliError("bad_args",
                           f"--freeze 需单个单元格,如 A2(冻结首行)/B1(冻结A列)/B2;"
                           f"收到区域 {args.freeze}")
        ws.freeze_panes = xlutil.area_label(r1, c1, r2, c2)
        applied.append(f"freeze={ws.freeze_panes}")

    # ---- 自动筛选 ----
    if args.unfilter:
        ws.auto_filter.ref = None
        applied.append("unfilter")
    elif args.filter is not None:
        r1, c1, r2, c2 = xlutil.parse_range(args.filter, ws)
        if r2 < r1 + 1:
            raise CliError("bad_args",
                           f"--filter 区域 {args.filter} 只有一行,须包含表头+数据")
        ws.auto_filter.ref = xlutil.area_label(r1, c1, r2, c2)
        applied.append(f"filter={ws.auto_filter.ref}")

    # ---- 隐藏 / 取消隐藏 ----
    if args.hide_cols is not None:
        c1, c2 = _plain_range(args.hide_cols, axis="col")
        for ci in range(c1, c2 + 1):
            ws.column_dimensions[get_column_letter(ci)].hidden = True
        applied.append(f"hide-cols {args.hide_cols}")
    if args.show_cols is not None:
        c1, c2 = _plain_range(args.show_cols, axis="col")
        for ci in range(c1, c2 + 1):
            ws.column_dimensions[get_column_letter(ci)].hidden = False
        applied.append(f"show-cols {args.show_cols}")
    if args.hide_rows is not None:
        r1, r2 = _plain_range(args.hide_rows, axis="row")
        for ri in range(r1, r2 + 1):
            ws.row_dimensions[ri].hidden = True
        applied.append(f"hide-rows {args.hide_rows}")
    if args.show_rows is not None:
        r1, r2 = _plain_range(args.show_rows, axis="row")
        for ri in range(r1, r2 + 1):
            ws.row_dimensions[ri].hidden = False
        applied.append(f"show-rows {args.show_rows}")

    xlutil.save_workbook_atomic(wb, plan.write_path)
    _up = {"upgraded_from": plan.upgraded_from} if plan.upgraded_from else {}
    return {"ok": True, "file": plan.write_path, **_up, "sheet": ws.title,
            "applied": applied}


# ---------------------------------------------------------------------------
# 规格解析
# ---------------------------------------------------------------------------

def _parse_specs(spec: str, *, axis: str):
    """'A=18,C:E=22' -> [(1,1,18.0),(3,5,22.0)];axis='row' 时 '1=28,3:5=20'。"""
    out = []
    for part in spec.split(","):
        part = part.strip()
        if "=" not in part:
            raise CliError("bad_args", f"规格 '{part}' 缺少 '=',应为 范围=数值")
        rng, _, val = part.partition("=")
        try:
            value = float(val.strip())
        except ValueError:
            raise CliError("bad_args", f"数值 '{val.strip()}' 不是数字") from None
        if value <= 0:
            raise CliError("bad_args", f"数值须为正数,收到 '{val.strip()}'")
        r1, r2 = _plain_range(rng.strip(), axis=axis)
        out.append((r1, r2, value))
    return out


def _plain_range(text: str, *, axis: str) -> tuple[int, int]:
    """列 'C'/'C:E' 或行 '1'/'3:5' -> (起, 止) 1-based;字母统一大写。"""
    t = text.strip().upper()
    if axis == "col":
        if re.fullmatch(r"[A-Z]+", t):
            return xlutil.col_to_idx(t), xlutil.col_to_idx(t)
        m = re.fullmatch(r"([A-Z]+):([A-Z]+)", t)
        if m and xlutil.col_to_idx(m.group(1)) <= xlutil.col_to_idx(m.group(2)):
            return xlutil.col_to_idx(m.group(1)), xlutil.col_to_idx(m.group(2))
        raise CliError("bad_args", f"列范围 '{text}' 格式不对,应为 C 或 C:E")
    m = re.fullmatch(r"(\d+)(?::(\d+))?", t)
    if m:
        r1 = int(m.group(1))
        r2 = int(m.group(2) or r1)
        if r1 >= 1 and r1 <= r2:
            return r1, r2
    raise CliError("bad_args", f"行范围 '{text}' 格式不对,应为 1 或 3:5")
