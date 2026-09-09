"""cond-format — 单元格条件格式(数值比较/文本包含/重复值高亮)。"""

from __future__ import annotations

import argparse
import re

from openpyxl.formatting.rule import CellIsRule, Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle

from .. import ioplan, xlutil
from ..cli import add_file_arg, add_sheet_arg
from ..errors import CliError

NAME = "cond-format"
HELP = "条件格式:add 加规则(大于/小于/区间/包含/重复值)/ clear 清除"
DESCRIPTION = """单元格条件格式(打开 Excel 时按值动态变色)。

用法示例:
  office cond-format add -f f.xlsx --range A2:A100 --op gt --value 100
  office cond-format add -f f.xlsx --range B2:B100 --op between --value '50,100' --fill C6EFCE
  office cond-format add -f f.xlsx --range C2:C100 --op contains --value '延期' --fill FFC7CE
  office cond-format add -f f.xlsx --range D2:D100 --op duplicates       # 重复值高亮
  office cond-format clear -f f.xlsx --range A2:A100    # 清除该区域规则
  office cond-format clear -f f.xlsx --all              # 清除整表规则

--op 取值: gt / lt / gte / lte / eq / neq / between / not-between / contains / duplicates
- between/not-between 的 --value 用逗号写两个边界,如 '50,100'
- contains 匹配文本片段(不区分大小写)
- --fill 默认 FFC7CE(浅红),--font-color 默认 9C0006(深红);可传 6 位 hex 自定义
- 同一区域多次 add 会叠加多条规则;clear 按区域整体移除
输出 JSON: {"ok": true, "file": "...", "rules": N}
"""

_NUM_RE = re.compile(r"^[+-]?\d+(?:\.\d+)?$")

_OPS = {
    "gt": "greaterThan", "lt": "lessThan", "gte": "greaterThanOrEqual",
    "lte": "lessThanOrEqual", "eq": "equal", "neq": "notEqual",
    "between": "between", "not-between": "notBetween",
}


def register(sp: argparse.ArgumentParser) -> None:
    sub = sp.add_subparsers(dest="cf_action", metavar="动作", required=True)

    s = sub.add_parser("add", help="添加一条条件格式规则")
    add_file_arg(s)
    add_sheet_arg(s)
    s.add_argument("--range", required=True, metavar="REF",
                   help="应用区域,如 A2:A100")
    s.add_argument("--op", required=True, choices=list(_OPS) + ["contains", "duplicates"],
                   help="规则类型(见上)")
    s.add_argument("--value", metavar="V",
                   help="阈值/文本;between 用 '50,100';duplicates 不需要")
    s.add_argument("--fill", metavar="HEX", default="FFC7CE",
                   help="命中底色(6 位 hex,默认 FFC7CE)")
    s.add_argument("--font-color", metavar="HEX", default="9C0006",
                   help="命中字体色(6 位 hex,默认 9C0006)")
    s.set_defaults(cf_func="add")

    s2 = sub.add_parser("clear", help="清除区域/全表条件格式")
    add_file_arg(s2)
    add_sheet_arg(s2)
    g = s2.add_mutually_exclusive_group(required=True)
    g.add_argument("--range", metavar="REF", help="要清除的区域")
    g.add_argument("--all", action="store_true", help="清除整表所有条件格式")
    s2.set_defaults(cf_func="clear")


def run(args: argparse.Namespace) -> dict:
    plan = ioplan.excel_plan(args.file)
    wb = xlutil.open_workbook(plan.read_path)
    ws = xlutil.choose_sheet(wb, args.sheet)

    if args.cf_func == "add":
        _add_rule(ws, args)
    else:
        _clear_rules(ws, args)

    xlutil.save_workbook_atomic(wb, plan.write_path)
    n = len(ws.conditional_formatting)
    _up = {"upgraded_from": plan.upgraded_from} if plan.upgraded_from else {}
    return {"ok": True, "file": plan.write_path, **_up, "sheet": ws.title,
            "rules": n}


# ---------------------------------------------------------------------------

def _fill(h: str) -> PatternFill:
    try:
        return PatternFill("solid", fgColor=_hex(h))
    except ValueError:
        raise CliError("bad_args", f"--fill 需 6 位 hex 颜色,收到 '{h}'") from None


def _hex(h: str) -> str:
    h = h.strip().lstrip("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}", h):
        raise CliError("bad_args", f"颜色需 6 位 hex(如 FF0000),收到 '{h}'")
    return h


def _formula(v: str) -> str:
    """数值按原样、其余加双引号(Excel 字符串字面量)。"""
    v = v.strip()
    return v if _NUM_RE.match(v) else f'"{v}"'


def _add_rule(ws, args) -> None:
    r1, c1, r2, c2 = xlutil.parse_range(args.range, ws)
    rng = xlutil.area_label(r1, c1, r2, c2)
    anchor = xlutil.area_label(r1, c1, r1, c1)

    fill = _fill(args.fill)
    font = Font(color=_hex(args.font_color))
    op = args.op

    if op in _OPS:
        values = []
        if op in ("between", "not-between"):
            if not args.value or "," not in args.value:
                raise CliError("bad_args",
                               f"--op {op} 需要 --value '下限,上限',如 '50,100'")
            lo, _, hi = args.value.partition(",")
            values = [_formula(lo), _formula(hi)]
        else:
            if args.value is None:
                raise CliError("bad_args", f"--op {op} 需要 --value")
            values = [_formula(args.value)]
        ws.conditional_formatting.add(
            rng, CellIsRule(operator=_OPS[op], formula=values,
                            fill=fill, font=font))
        return

    dxf = DifferentialStyle(fill=fill, font=font)
    if op == "contains":
        if not args.value:
            raise CliError("bad_args", "--op contains 需要 --value 文本")
        text = args.value
        ws.conditional_formatting.add(
            rng, Rule(type="containsText", operator="containsText", text=text,
                      formula=[f'NOT(ISERROR(SEARCH("{text}",{anchor})))'],
                      dxf=dxf))
        return
    if op == "duplicates":
        ws.conditional_formatting.add(rng, Rule(type="duplicateValues", dxf=dxf))
        return
    raise CliError("internal", f"未知规则 {op}")  # pragma: no cover


def _clear_rules(ws, args) -> None:
    if args.all:
        cf = ws.conditional_formatting
        keys = [str(e.sqref) for e in cf]
        for k in keys:
            try:
                del cf[k]
            except KeyError:
                pass
        return
    r1, c1, r2, c2 = xlutil.parse_range(args.range, ws)
    key = xlutil.area_label(r1, c1, r2, c2)
    try:
        del ws.conditional_formatting[key]
    except KeyError:
        pass  # 该区域本来就没有规则,视为已清除
