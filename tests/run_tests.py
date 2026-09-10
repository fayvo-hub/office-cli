"""office-cli 端到端测试:每个用例用 subprocess 调用真实 CLI,校验退出码与 JSON 输出。

运行: python tests/run_tests.py   (在项目根目录)
输出: 每行 PASS/FAIL,最后汇总;失败可加环境变量 VERBOSE=1 打印输出详情。
"""

from __future__ import annotations

import csv
import datetime as _dt
import io
import json
import os
import shutil
import subprocess
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLES = os.path.join(ROOT, "samples")
WORK = os.path.join(ROOT, ".test-data")

# excel 组的命令名(旧测试直接以命令名开头,自动补组前缀)
_EXCEL_CMDS = {"list", "read", "write", "sheet", "style", "merge",
               "chart", "image", "pivot", "layout", "cond-format",
               "comment", "validate", "insert", "delete", "replace"}

VERBOSE = os.environ.get("VERBOSE") == "1"
_passed = 0
_failed = 0
_failures = []


def run(*argv: str, env: dict | None = None) -> tuple[int, dict | None, dict | None]:
    """调用 office;返回 (退出码, stdout JSON, stderr JSON)。

    excel 组命令(list/read/write/...)自动加 "excel" 前缀(兼容旧式调用)。
    env: 额外环境变量(如 OFFICE_ENGINE=lo 测跨平台引擎),与当前环境合并。
    """
    if argv and argv[0] in _EXCEL_CMDS:
        argv = ("excel",) + argv
    proc = subprocess.run(
        [sys.executable, "-m", "office", *argv],
        capture_output=True, text=True, encoding="utf-8",
        cwd=ROOT, timeout=180, env={**os.environ, **(env or {})},
    )
    out = err = None
    try:
        out = json.loads(proc.stdout) if proc.stdout.strip() else None
    except json.JSONDecodeError:
        out = {"__raw__": proc.stdout}
    try:
        err = json.loads(proc.stderr) if proc.stderr.strip() else None
    except json.JSONDecodeError:
        err = {"__raw__": proc.stderr}
    return proc.returncode, out, err


def check(name: str, cond: bool, detail: str = "") -> None:
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"PASS  {name}")
    else:
        _failed += 1
        _failures.append(name)
        print(f"FAIL  {name}  {detail[:400]}")


def expect_ok(name, *argv, env=None, **kw):
    code, out, err = run(*argv, env=env)
    ok = code == 0 and out is not None and out.get("ok") is True
    detail = f"code={code} out={out} err={err}" if not ok or VERBOSE else ""
    check(name, ok, detail)
    return out


def expect_err(name, errcode, *argv, env=None):
    code, out, err = run(*argv, env=env)
    got = (err is not None and err.get("error", {}).get("code") == errcode
           and code != 0)
    detail = f"code={code} err={err}" if not got or VERBOSE else ""
    check(name, got, detail)


def run_py(code: str, env: dict | None = None) -> tuple[int, str, str]:
    """跑一段内联 Python(直接导入 office 包做单元级验证);返回 (码, stdout, stderr)。"""
    proc = subprocess.run(
        [sys.executable, "-X", "utf8", "-c", code],
        capture_output=True, text=True, encoding="utf-8", cwd=ROOT, timeout=180,
        env={**os.environ, **(env or {})},
    )
    return proc.returncode, (proc.stdout or "").strip(), (proc.stderr or "").strip()


def py_json(code: str, env: dict | None = None):
    """跑内联 Python 并把最后一行 stdout 当 JSON 解析(失败返回 None)。"""
    rc, out, _err = run_py(code, env)
    if rc != 0 or not out:
        return None
    try:
        return json.loads(out.splitlines()[-1])
    except json.JSONDecodeError:
        return None


def wpath(name: str) -> str:
    """工作副本路径(测试用,随便改)"""
    return os.path.join(WORK, name)


def copy_sample(name: str) -> str:
    shutil.copy(os.path.join(SAMPLES, name), wpath(name))
    return wpath(name)


# ---------------------------------------------------------------------------

def display_checks() -> None:
    """渲染层单测: number_format → 输出形态(真实报表带 [Red] 颜色标记, 曾误判为日期)。"""
    from office.cmds_rag.prep import _display
    ck = lambda name, cond: check(f"display {name}", cond, str((cond, _display)))
    # [Red] 颜色标记里的字母不得命中日期占位符(5900 曾变 1916-02-25)
    ck("[Red]金额不误判日期", _display(5900, "#,##0;[Red]-#,##0") == 5900)
    ck("[Red]货币不误判日期",
       _display(862000, "¥#,##0.00;[Red]-¥#,##0.00") == 862000)
    # 真日期格式: 裸数值按序列转 ISO; 带时间才出时间
    ck("日期序列→ISO", _display(45292, "yyyy-mm-dd") == "2024-01-01")
    ck("日期时间保留", _display(45292.5, "yyyy-mm-dd hh:mm")
       == "2024-01-01T12:00:00")
    ck("剥色后仍识别日期", _display(5900, "[Red]yyyy-mm-dd")
       == "1916-02-25")
    ck("百分比渲染", _display(0.13, "0.0%") == "13%")
    ck("文本原样", _display("苹果", "yyyy") == "苹果")


def formula_checks() -> None:
    """RAG 公式求值器单测(内存环境直测,不经 CLI)。"""
    import datetime as dt2

    from office.cmds_rag.formulas import FormulaError, evaluate

    class Mem:
        TABLES = {}
        stack = []

        def __init__(self):
            self.sales = {
                (1, 1): "苹果", (2, 1): "香蕉", (3, 1): "苹果",
                (4, 1): "梨", (5, 1): "橙子",
                (1, 2): 100, (2, 2): 150, (3, 2): 200, (4, 2): 80, (5, 2): 60,
                (1, 3): 10.0, (2, 3): 5.5, (3, 3): 12.0, (4, 3): 8.0, (5, 3): 20.0,
                (1, 4): "=B1*C1", (2, 4): "=B2*C2", (3, 4): "=B3*C3",
                (4, 4): "=B4*C4", (5, 4): "=B5*C5",
                (1, 5): "", (2, 5): "好", (3, 5): "", (4, 5): "", (5, 5): "差",
                (7, 1): "苹果", (7, 2): 300, (7, 3): 9.0, (7, 4): "=B7*C7",
                (9, 1): "=B9", (9, 2): "=A9",
                (1, 6): 1, (2, 6): 2, (3, 6): 3, (4, 6): 4, (5, 6): 5,
                (1, 7): "一", (2, 7): "二", (3, 7): "三", (4, 7): "四", (5, 7): "五",
                (1, 8): 100.0, (2, 8): -50.0, (3, 8): -40.0, (4, 8): -30.0,
            }
            self.TABLES = {
                "销售": self.sales,
                "汇总": {(1, 1): "标题", (1, 2): 1, (2, 1): 2, (2, 2): 3},
                "人 员表": {(1, 1): "姓名", (1, 2): "城市", (2, 1): "张三",
                            (2, 2): "北京", (3, 1): 5, (3, 2): "上海"},
            }
            self.stack = []

        def cell(self, sheet, row, col):
            t = self.TABLES.get(sheet or "销售")
            if t is None:
                raise FormulaError(f"引用的工作表不存在: {sheet}")
            v = t.get((row, col))
            if isinstance(v, str) and v.startswith("="):
                if (sheet, row, col) in self.stack:
                    raise FormulaError("循环引用")
                self.stack.append((sheet, row, col))
                try:
                    v = evaluate(self, sheet, row, col, v)
                finally:
                    self.stack.pop()
            return v

        def sheet_max_row(self, sheet):
            t = self.TABLES.get(sheet or "销售")
            return max((r for (r, c) in t), default=0) if t else 1048576

        def sheet_exists(self, sheet):
            return (sheet or "销售") in self.TABLES

    env = Mem()

    def ev(f):
        return evaluate(env, "销售", 1, 1, f)

    def evr(f):
        try:
            ev(f)
            return False
        except FormulaError:
            return True

    def ck(name, cond, detail=""):
        check(f"formula {name}", cond, detail)

    ck("加法", ev("=1+2") == 3)
    ck("括号", ev("=(1+2)*3") == 9)
    ck("幂", ev("=2^10") == 1024)
    ck("负号", ev("=-5+2") == -3)
    ck("百分比", ev("=50%") == 0.5)
    ck("文本拼接", ev('="a"&"b"') == "ab")
    ck("除法", ev("=10/4") == 2.5)
    ck("除零", evr("=1/0"))
    ck("单格取值", ev("=B1") == 100)
    ck("算术引用", ev("=B1*C1") == 1000.0)
    ck("负引用", ev("=-B2") == -150)
    ck("引用拼接", ev('=E2&"-"&B2') == "好-150")
    ck("引用比较", ev('=IF(A1="苹果","是","否")') == "是")
    ck("引用嵌套求和", ev("=SUM(B1:C1)+SUM(B3:C3)") == 322.0)
    ck("百分比引用", ev("=C1%") == 0.1)
    ck("幂引用", ev("=B1^2") == 10000.0)
    ck("顶层裸引用", ev("=B2") == 150)
    ck("空引用算术", ev("=E1+1") == 1)
    ck("&空引用", ev('=A1&E1&"!"') == "苹果!")
    ck("LEFT引用", ev("=LEFT(A1,1)") == "苹")
    ck("RIGHT引用", ev("=RIGHT(A1,1)") == "果")
    ck("MID引用", ev("=MID(A1,2,1)") == "果")
    ck("连接截取", ev("=LEFT(A1,1)&RIGHT(A1,1)&MID(A1,2,1)") == "苹果果")
    ck("LEN", ev("=LEN(A1&B2)") == 5)
    ck("SUBSTITUTE", ev('=SUBSTITUTE(A1,"苹果","橙子")') == "橙子")
    ck("TRIM", ev('=TRIM("  a  b ")') == "a b")
    ck("UPPER", ev('=UPPER("ab")') == "AB")
    ck("TODAY类型", isinstance(ev("=TODAY()"), dt2.date))
    ck("NOW类型", isinstance(ev("=NOW()"), dt2.datetime))
    ck("DATE构造", ev("=DATE(2024,3,15)") == _dt.date(2024, 3, 15))
    ck("DATE引用参数", ev("=DATE(2024,C1,15)") == _dt.date(2024, 10, 15))
    ck("DATE非法", evr("=DATE(2024,13,1)"))
    ck("裸表名取值", ev("=汇总!A1") == "标题")
    ck("裸表名算术", ev("=汇总!B2*2") == 6)
    ck("跨表单格SUM忽略文本", ev("=SUM(汇总!A1)") == 0)
    ck("SUMIF整列跨表", ev("=SUMIF(汇总!B1:B2,1)") == 1)
    ck("跨表区域SUM", ev("=SUM(汇总!B1:B2)") == 4.0)
    ck("引号表名取值", ev("='人 员表'!B2") == "北京")
    ck("引号表名算术", ev("='人 员表'!A3+1") == 6)
    ck("引号表名SUM区域", ev("=SUM('人 员表'!A2:A3)") == 5.0)
    ck("引号文本字面量", ev("=A1='苹果'") is True)
    ck("sq无感叹号当文本", ev('="前缀"&\'abc\'') == "前缀abc")
    ck("SUM区域", ev("=SUM(B1:B5)") == 590.0)
    ck("SUM多区", ev("=SUM(B1:B2,C1:C2)") == 265.5)
    ck("AVERAGE", ev("=AVERAGE(B1:B5)") == 118.0)
    ck("COUNT", ev("=COUNT(B1:B5)") == 5)
    ck("COUNTBLANK", ev("=COUNTBLANK(E1:E5)") == 3)
    ck("COUNTA", ev("=COUNTA(E1:E5)") == 2)
    ck("MAX/MIN", ev("=MAX(B1:B5)") == 200 and ev("=MIN(B1:B5)") == 60)
    ck("MAX空区0", ev("=MAX(A20:A21)") == 0)
    ck("SUM空文本忽略", ev("=SUM(A1:A5)") == 0)
    ck("SUMIF文本相等", ev('=SUMIF(A1:A5,"苹果",B1:B5)') == 300)
    ck("SUMIF通配", ev('=SUMIF(A1:A5,"苹*",B1:B5)') == 300)
    ck("SUMIF数值", ev("=SUMIF(B1:B5,200)") == 200)
    ck("SUMIF比较", ev('=SUMIF(B1:B5,">=150",C1:C5)') == 17.5)
    ck("SUMIF单格扩展", ev('=SUMIF(A1:A5,"苹果",B1)') == 300)
    ck("SUMIF无sum_range", ev('=SUMIF(B1:B5,">100")') == 350)
    ck("SUMIFS多条件", ev('=SUMIFS(B1:B5,A1:A5,"苹果",C1:C5,">9")') == 300)
    ck("COUNTIF", ev('=COUNTIF(A1:A5,"苹果")') == 2)
    ck("COUNTIFS", ev('=COUNTIFS(A1:A5,"苹果",C1:C5,">9")') == 2)
    ck("SUMIF文本匹配区", ev('=SUMIF(A1:A5,"苹果")') == 0)
    ck("IF空引用条件", ev('=IF(E1,"有","无")') == "无")
    ck("AND/OR/NOT", ev("=AND(B1>50,B2>50)") is True
       and ev("=NOT(1>2)") is True)
    ck("ISNUMBER", ev("=ISNUMBER(B1)") is True
       and ev("=ISNUMBER(A1)") is False
       and ev("=ISNUMBER(E1)") is False
       and ev("=ISNUMBER(TRUE)") is False
       and ev("=ISNUMBER(DATE(2024,1,1))") is True)
    ck("ISBLANK", ev("=ISBLANK(A6)") is True
       and ev("=ISBLANK(B1)") is False
       and ev('=ISBLANK("")') is False)  # Excel: 空串不是 blank
    ck("ISTEXT/ISNONTEXT", ev("=ISTEXT(A1)") is True
       and ev("=ISNONTEXT(B1)") is True
       and ev("=ISNONTEXT(A1)") is False)
    ck("ISLOGICAL", ev("=ISLOGICAL(1>2)") is True)
    ck("IF+ISNUMBER 级联", ev('=IF(ISNUMBER(B1),"数值","非数值")') == "数值"
       and ev('=IF(ISNUMBER(A1),"数值","非数值")') == "非数值")
    ck("空区AVERAGE", evr("=AVERAGE(A20:A21)"))
    ck("未知函数", evr("=VLOOKUP(1,2,3)"))
    ck("不存在表", evr("='不存在'!A1+1"))
    ck("裸区域运算", evr("=B1:B3+1"))
    ck("多格参数", evr("=LEFT(A1:B2,1)"))
    ck("未闭合括号", evr("=SUM(A1:B2"))
    ck("循环引用", evr("=A9"))
    ck("文本算术", evr('="abc"+1'))
    ck("TRUE/FALSE", ev("=TRUE") is True and ev("=FALSE") is False)
    ck("ROUND", ev("=ROUND(B1/3,2)") == 33.33)
    # ---- v3 新函数: 惰性/错误捕获 ----
    ck("IF惰性分支", ev('=IF(B1>50,"ok",1/0)') == "ok")
    ck("IFERROR除零", ev('=IFERROR(1/0,"兜底")') == "兜底")
    ck("IFERROR正常", ev('=IFERROR(B1,"兜底")') == 100)
    ck("IFERROR包查找", ev('=IFERROR(VLOOKUP("柚子",A1:B5,2,FALSE),"无")') == "无")
    ck("IFNA未命中", ev('=IFNA(VLOOKUP("柚子",A1:B5,2,FALSE),"无")') == "无")
    ck("ISERROR/ISNA/ISERR", ev("=ISERROR(1/0)") is True
       and ev("=ISERROR(B1)") is False
       and ev("=ISNA(NA())") is True
       and ev("=ISERR(NA())") is False
       and ev("=ISERR(1/0)") is True)
    ck("IFS", ev('=IFS(B1>200,"高",B1>50,"中",TRUE,"低")') == "中")
    ck("SWITCH", ev('=SWITCH(A1,"苹果","A","香蕉","B","其他")') == "A")
    ck("CHOOSE", ev('=CHOOSE(2,"一","二","三")') == "二")
    ck("CHOOSE越界", evr('=CHOOSE(5,"一")'))
    ck("NA错误", evr("=NA()"))
    # ---- v3 新函数: 查找族 ----
    ck("VLOOKUP精确", ev('=VLOOKUP("香蕉",A1:B5,2,FALSE)') == 150)
    ck("VLOOKUP未命中", evr('=VLOOKUP("柚子",A1:B5,2,FALSE)'))
    ck("VLOOKUP列越界", evr('=VLOOKUP("苹果",A1:B5,9,FALSE)'))
    ck("VLOOKUP近似升序", ev("=VLOOKUP(3.5,F1:G5,2,TRUE)") == "三")
    ck("VLOOKUP通配", ev('=VLOOKUP("苹*",A1:B5,2,FALSE)') == 100)
    ck("HLOOKUP", ev("=HLOOKUP(100,A1:C5,3,FALSE)") == 200)
    ck("MATCH精确/近似/通配", ev('=MATCH("橙子",A1:A5,0)') == 5
       and ev("=MATCH(3.5,F1:F5,1)") == 3
       and ev('=MATCH("梨*",A1:A5,0)') == 4)
    ck("MATCH未命中", evr('=MATCH("柚子",A1:A5,0)'))
    ck("INDEX行列", ev("=INDEX(A1:C5,2,2)") == 150)
    ck("INDEX单列", ev("=INDEX(B1:B5,3)") == 200)
    ck("INDEX整列引用", ev("=SUM(INDEX(A1:C5,0,2))") == 590.0)
    ck("INDEX越界", evr("=INDEX(A1:C5,9,2)"))
    ck("XLOOKUP", ev('=XLOOKUP("橙子",A1:A5,B1:B5)') == 60)
    ck("XLOOKUP默认值", ev('=XLOOKUP("柚子",A1:A5,B1:B5,"无")') == "无")
    ck("XLOOKUP未找到报错", evr('=XLOOKUP("柚子",A1:A5,B1:B5)'))
    ck("XLOOKUP向下/向上近似", ev("=XLOOKUP(3.5,F1:F5,G1:G5,0,-1)") == "三"
       and ev("=XLOOKUP(3.5,F1:F5,G1:G5,0,1)") == "四")
    ck("LOOKUP近似", ev("=LOOKUP(3.5,F1:F5,G1:G5)") == "三")
    # ---- v3 新函数: 统计/条件聚合 ----
    ck("SUMPRODUCT", ev("=SUMPRODUCT(B1:B5,C1:C5)") == 6065.0)
    ck("MEDIAN/LARGE/SMALL", ev("=MEDIAN(B1:B5)") == 100
       and ev("=LARGE(B1:B5,2)") == 150 and ev("=SMALL(B1:B5,2)") == 80)
    ck("RANK", ev("=RANK(150,B1:B5)") == 2)
    ck("STDEV/VAR", ev("=ROUND(STDEV(B1:B5),2)") == 56.75
       and ev("=VAR.P(B1:B5)") == 2576.0)
    ck("PRODUCT", ev("=PRODUCT(F1:F5)") == 120)
    ck("PERCENTILE", ev("=PERCENTILE(B1:B5,0.5)") == 100)
    ck("COUNTIF不等于", ev('=COUNTIF(A1:A5,"<>苹果")') == 3)
    ck("AVERAGEIF", ev('=AVERAGEIF(A1:A5,"苹果",B1:B5)') == 150)
    ck("AVERAGEIFS", ev('=ROUND(AVERAGEIFS(B1:B5,A1:A5,"<>苹果"),2)') == 96.67)
    ck("MAXIFS/MINIFS", ev('=MAXIFS(B1:B5,A1:A5,"<>苹果")') == 150
       and ev('=MINIFS(B1:B5,A1:A5,"<>苹果")') == 60)
    # ---- v3 新函数: 数学/财务 ----
    ck("SQRT/LOG/LOG10", ev("=SQRT(16)") == 4 and ev("=LOG(8,2)") == 3
       and ev("=LOG10(1000)") == 3)
    ck("POWER/EXP", ev("=POWER(2,10)") == 1024
       and ev("=ROUND(EXP(1),5)") == 2.71828)
    ck("CEILING/FLOOR/MROUND", ev("=CEILING(2.1,1)") == 3
       and ev("=FLOOR(2.9,1)") == 2 and ev("=MROUND(2.5,1)") == 3)
    ck("TRUNC/SIGN", ev("=TRUNC(3.14159,2)") == 3.14 and ev("=SIGN(-5)") == -1)
    ck("EVEN/ODD/FACT", ev("=EVEN(3)") == 4 and ev("=ODD(2)") == 3
       and ev("=FACT(5)") == 120)
    ck("SUMSQ", ev("=SUMSQ(B1:B5)") == 82500)
    ck("PMT/FV/PV", ev("=ROUND(PMT(0.05/12,120,200000),2)") == -2121.31
       and ev("=ROUND(FV(0.05,10,-1000),2)") == 12577.89
       and ev("=ROUND(PV(0.05,10,-1000),2)") == 7721.73)
    ck("NPV", ev("=ROUND(NPV(0.1,B1:B3),2)") == 365.14)
    ck("IRR", ev("=ROUND(IRR(H1:H4),4)") == 0.1065)
    # ---- v3 新函数: 引用族 ----
    ck("ROW/COLUMN", ev("=ROW()") == 1 and ev("=COLUMN()") == 1
       and ev("=ROW(B3)") == 3 and ev("=COLUMN(C1)") == 3)
    ck("ROWS/COLUMNS", ev("=ROWS(B1:B5)") == 5 and ev("=COLUMNS(A1:C1)") == 3)
    ck("OFFSET区域", ev("=SUM(OFFSET(B1,1,1,2,1))") == 17.5)
    ck("OFFSET单格", ev("=OFFSET(B1,4,0)") == 60)
    ck("OFFSET越界", evr("=OFFSET(A1,-1,0)"))
    # ---- v3 新函数: 日期族 ----
    ck("YEAR/MONTH/DAY", ev("=YEAR(DATE(2024,1,12))") == 2024
       and ev("=MONTH(DATE(2024,1,12))") == 1 and ev("=DAY(DATE(2024,1,12))") == 12)
    ck("EOMONTH/EDATE", ev("=DAY(EOMONTH(DATE(2024,1,15),1))") == 29
       and ev("=DAY(EDATE(DATE(2024,1,31),1))") == 29)
    ck("DATEDIF", ev('=DATEDIF(DATE(2020,1,1),DATE(2024,3,5),"Y")') == 4
       and ev('=DATEDIF(DATE(2020,1,1),DATE(2024,3,5),"M")') == 50
       and ev('=DATEDIF(DATE(2020,1,1),DATE(2024,3,5),"D")') == 1525)
    ck("DAYS/WEEKDAY", ev("=DAYS(DATE(2024,1,12),DATE(2024,1,1))") == 11
       and ev("=WEEKDAY(DATE(2024,1,12))") == 6
       and ev("=WEEKDAY(DATE(2024,1,12),2)") == 5)
    ck("NETWORKDAYS/WORKDAY", ev("=NETWORKDAYS(DATE(2024,1,1),DATE(2024,1,12))") == 10
       and ev("=NETWORKDAYS(DATE(2024,1,1),DATE(2024,1,12),DATE(2024,1,3))") == 9
       and ev("=DAY(WORKDAY(DATE(2024,1,12),1))") == 15)
    ck("DATEVALUE", ev('=DATEVALUE("2024年1月12日")') == _dt.date(2024, 1, 12)
       and ev('=DATEVALUE("2024-01-12")') == _dt.date(2024, 1, 12))
    ck("TIME/HOUR/MINUTE", ev("=HOUR(TIME(13,30,0))") == 13
       and ev("=MINUTE(TIME(13,30,0))") == 30)
    # ---- v3 新函数: 文本族 ----
    ck("TEXT百分比", ev('=TEXT(0.135,"0.0%")') == "13.5%")
    ck("TEXT千分位", ev('=TEXT(1234.5,"#,##0.00")') == "1,234.50")
    ck("TEXT日期", ev('=TEXT(DATE(2024,1,12),"yyyy年m月d日")') == "2024年1月12日")
    ck("TEXTJOIN/CONCAT", ev('=TEXTJOIN("-",TRUE,"a","","b")') == "a-b"
       and ev('=CONCAT(A1:A2)') == "苹果香蕉")
    ck("FIND/SEARCH", ev('=FIND("果",A1)') == 2 and ev('=SEARCH("A","abc")') == 1)
    ck("FIND未找到", evr('=FIND("柚子",A1)'))
    ck("REPLACE/REPT", ev('=REPLACE("2024年01月",6,2,"12")') == "2024年12月"
       and ev('=REPT("ab",3)') == "ababab")
    ck("EXACT", ev('=EXACT("a","A")') is False and ev('=EXACT("a","a")') is True)
    ck("VALUE", ev('=VALUE("¥1,234.50")') == 1234.5 and ev('=VALUE("15%")') == 0.15)
    ck("PROPER/CLEAN", ev('=PROPER("hello world")') == "Hello World"
       and ev('=CLEAN("a"&CHAR(9)&"b")') == "ab")
    ck("CHAR/CODE", ev("=CHAR(65)") == "A" and ev('=CODE("A")') == 65)


def main() -> int:
    global _passed, _failed
    shutil.rmtree(WORK, ignore_errors=True)
    os.makedirs(WORK)
    demo = copy_sample("demo.xlsx")
    cached = copy_sample("demo_cached.xlsx")
    pivot_f = copy_sample("demo_pivot.xlsx")
    sample_md = copy_sample("sample.md")
    sample_docx = copy_sample("sample.docx")
    sample_pptx = copy_sample("sample.pptx")
    legacy_xls = os.path.join(SAMPLES, "legacy.xls")
    legacy_doc = os.path.join(SAMPLES, "legacy.doc")
    legacy_ppt = os.path.join(SAMPLES, "legacy.ppt")

    # ---------- list ----------
    out = expect_ok("list 基本", "list", "-f", demo)
    if out:
        names = [s["name"] for s in out["sheets"]]
        check("list 三个表", names == ["销售", "人员", "空表"], str(names))
        s1 = out["sheets"][0]
        check("list 销售表信息", s1["max_row"] == 16 and s1["max_column"] == 5
              and "A16:B16" in s1["merged_cells"] and s1["charts"] == 0,
              json.dumps(s1, ensure_ascii=False)[:300])
        check("list 预览含表头", s1["preview_rows"][0][0] == "月份",
              str(s1["preview_rows"][:1]))
    expect_err("list 文件不存在", "no_file", "list", "-f", wpath("nope.xlsx"))

    # ---------- read ----------
    out = expect_ok("read 全表(默认激活表)", "read", "-f", demo)
    check("read rows=16 行", out and out["row_count"] == 16, str(out))
    check("read D2 是公式文本", out and out["rows"][1][3] == "=B2*C2",
          str(out and out["rows"][1]))
    check("read 日期是 ISO", out and out["rows"][1][4].startswith("2024-01-15"),
          str(out and out["rows"][1][4]))
    check("read 合并区报告", out and "A16:B16" in out["merged_cells"],
          str(out and out["merged_cells"]))

    out = expect_ok("read 指定表", "read", "-f", demo, "--sheet", "人员")
    check("read 人员中文数据", out and out["rows"][1][0] == "张三"
          and out["rows"][1][1] == "北京", str(out and out["rows"][1]))

    out = expect_ok("read 单格", "read", "-f", demo, "--range", "B2")
    check("read 单格值", out and out["rows"] == [[100]], str(out and out["rows"]))

    out = expect_ok("read 区域 + cached", "read", "-f", cached, "--range", "A1:D4",
                    "--cached")
    check("read cached 公式=计算值", out and out["rows"][1][3] == 990.0,
          str(out and out["rows"][1]))

    out = expect_ok("read cells 模式", "read", "-f", cached, "--range", "D2",
                    "--cells", "--cached")
    c = out["cells"][0] if out else {}
    check("read cells 元数据", c.get("type") == "formula" and c.get("formula") == "=B2*C2"
          and c.get("value") == 990.0 and c.get("cached_available") is True,
          json.dumps(c, ensure_ascii=False))

    out = expect_ok("read cells 注释与超链接", "read", "-f", demo, "--range", "B2",
                    "--cells")
    c = out["cells"][0] if out else {}
    check("read comment/hyperlink", "comment" in c and c["comment"].startswith("1 月销量"),
          json.dumps(c, ensure_ascii=False)[:200])

    out = expect_ok("read 整列 + limit", "read", "-f", demo, "--range", "A:A",
                    "--limit", "3")
    check("read 截断警告", out and out["row_count"] == 3
          and any("limit" in w for w in out["warnings"]),
          str(out and (out["row_count"], out["warnings"])))

    out = expect_ok("read 空表", "read", "-f", demo, "--sheet", "空表")
    check("read 空表 rows=[]", out and out["rows"] == [], str(out))

    expect_err("read 表不存在", "no_sheet", "read", "-f", demo, "--sheet", "不存在")
    expect_err("read 坏范围", "bad_range", "read", "-f", demo, "--range", "A0")

    # ---------- write ----------
    out = expect_ok("write 单值", "write", "-f", demo, "--cell", "C15", "--data", "123.5")
    check("write 返回范围", out and out["range"] == "C15", str(out))
    out = expect_ok("write 后读回", "read", "-f", demo, "--range", "C15")
    check("write 值生效", out and out["rows"] == [[123.5]], str(out and out["rows"]))

    out = expect_ok("write 二维+公式", "write", "-f", demo, "--cell", "F1",
                    "--data", '[[1, "=F1*2"], [2, "文字"]]')
    check("write 范围 F1:G2", out and out["range"] == "F1:G2"
          and out["values_written"] == 4, str(out))
    out = expect_ok("write 公式读回", "read", "-f", demo, "--range", "G1", "--cells")
    check("write 公式格", out and out["cells"][0]["type"] == "formula", str(out))

    out = expect_ok("write --literal", "write", "-f", demo, "--cell", "F4",
                    "--data", '"=NOT_A_FORMULA"', "--literal")
    out = expect_ok("write literal 读回", "read", "-f", demo, "--range", "F4",
                    "--cells")
    check("write literal 是文本", out and out["cells"][0]["type"] == "str"
          and out["cells"][0]["value"] == "=NOT_A_FORMULA", json.dumps(out, ensure_ascii=False))

    jf = wpath("vals.json")
    with open(jf, "w", encoding="utf-8") as fh:
        json.dump([["甲", 1], ["乙", 2]], fh, ensure_ascii=False)
    expect_ok("write --data-file", "write", "-f", demo, "--cell", "H1",
              "--data-file", jf)

    created = wpath("created.xlsx")
    expect_ok("write --create 新建", "write", "-f", created, "--cell", "A1",
              "--data", "1", "--create")
    expect_err("write --create 撞已存在", "file_exists", "write", "-f", created,
               "--cell", "A1", "--data", "1", "--create")
    expect_ok("write --create --overwrite", "write", "-f", created, "--cell", "A1",
              "--data", '"x"', "--create", "--overwrite")
    expect_err("write 文件不存在(无 create)", "no_file", "write", "-f",
               wpath("missing.xlsx"), "--data", "1")
    expect_err("write 坏 JSON", "bad_json", "write", "-f", created, "--data", "{oops")
    expect_err("write 越界", "too_large", "write", "-f", created, "--cell", "A1048570",
               "--data", '[[1],[2],[3],[4],[5],[6],[7],[8],[9],[10]]')

    # ---------- sheet ----------
    out = expect_ok("sheet add", "sheet", "add", "-f", demo, "--name", "新表")
    check("sheet add 出现", out and "新表" in out["sheets"], str(out))
    expect_err("sheet add 重名", "dup_sheet", "sheet", "add", "-f", demo,
               "--name", "销售")
    expect_ok("sheet rename", "sheet", "rename", "-f", demo, "--name", "新表",
              "--new-name", "改后")
    out = expect_ok("sheet copy", "sheet", "copy", "-f", demo, "--name", "销售",
                    "--new-name", "销售副本")
    check("sheet copy 含副本", out and "销售副本" in out["sheets"], str(out))
    out = expect_ok("sheet remove", "sheet", "remove", "-f", demo, "--name", "改后")
    check("sheet remove 生效", out and "改后" not in out["sheets"], str(out))
    out = expect_ok("sheet remove 销售副本", "sheet", "remove", "-f", demo,
                    "--name", "销售副本")

    # active 用独立副本,避免改变 demo 的默认激活表影响后续用例
    active_sh = copy_sample("demo.xlsx")
    expect_ok("sheet active", "sheet", "active", "-f", active_sh, "--name", "人员")
    out = expect_ok("sheet active 后 list", "list", "-f", active_sh)
    wb_tmp = load_workbook(active_sh)
    check("sheet active 生效", wb_tmp.active.title == "人员", wb_tmp.active.title)

    # ---------- style ----------
    st = wpath("style.xlsx")
    shutil.copy(demo, st)
    out = expect_ok("style 表头", "style", "-f", st, "--sheet", "销售",
                    "--range", "A1:E1",
                    "--bold", "--fill", "FFC000", "--align", "center",
                    "--font-size", "13", "--font-color", "FFFFFF",
                    "--border", "thin", "--num-format", "0.0")
    check("style cells_affected=5", out and out["cells_affected"] == 5, str(out))
    wb = load_workbook(st)
    c = wb["销售"]["A1"]
    check("style 字体生效", c.font.bold and c.font.size == 13, str(c.font))
    check("style 填充生效", c.fill.fgColor.rgb in ("00FFC000", "FFFFC000"),
          str(c.fill.fgColor.rgb))
    check("style 数字格式生效", c.number_format == "0.0", c.number_format)
    expect_err("style 颜色非法", "bad_color", "style", "-f", st, "--range", "A1",
               "--fill", "red")
    expect_err("style 区域过大", "too_large", "style", "-f", st,
               "--range", "A1:C20000", "--bold")

    # ---------- merge ----------
    mg = wpath("merge.xlsx")
    shutil.copy(demo, mg)
    S = ["--sheet", "销售"]
    out = expect_ok("merge 合并", "merge", "-f", mg, *S, "--range", "C17:D17")
    check("merge 报告合并区", out and "C17:D17" in out["merged_cells"], str(out))
    # 向合并区非左上角写入 -> 业务错误 merged_cell(而非 internal)
    expect_err("write 合并区非左上格", "merged_cell", "write", "-f", mg, *S,
               "--cell", "D17", "--data", "1")
    # 合并区左上角仍可写(数据只在左上角)
    expect_ok("write 合并区左上角", "write", "-f", mg, *S, "--cell", "C17",
              "--data", '"更新"')
    expect_err("merge 重叠", "overlap", "merge", "-f", mg, *S, "--range", "D17:E17")
    out = expect_ok("write 准备合并警告", "write", "-f", mg, *S, "--cell", "A18",
                    "--data", '"保留"')
    expect_ok("write 准备合并警告2", "write", "-f", mg, *S, "--cell", "B18",
              "--data", '"丢弃我"')
    out = expect_ok("merge 非左上值警告", "merge", "-f", mg, *S, "--range", "A18:B18")
    check("merge 警告存在", out and any("B18" in w for w in out["warnings"]),
          str(out and out["warnings"]))
    out = expect_ok("merge unmerge", "merge", "-f", mg, *S, "--range", "C17:D17",
                    "--unmerge")
    check("merge unmerge 生效", out and "C17:D17" not in out["merged_cells"], str(out))
    expect_err("merge unmerge 非合并区", "not_merged", "merge", "-f", mg, *S,
               "--range", "E1:F1", "--unmerge")

    # ---------- convert ----------
    o_csv = wpath("out.csv")
    expect_ok("convert xlsx->csv", "convert", "-f", cached, "--sheet", "销售",
              "--out", o_csv, "--cached")
    with open(o_csv, encoding="utf-8-sig") as fh:
        lines = fh.readlines()
    check("convert csv BOM+公式值", lines[1].startswith("2024-01,")
          and "990.0" in lines[1], lines[1] if lines else "")

    o_json = wpath("out.json")
    out = expect_ok("convert xlsx->json", "convert", "-f", cached, "--sheet", "销售",
                    "--out", o_json, "--cached")
    with open(o_json, encoding="utf-8") as fh:
        doc = json.load(fh)
    check("convert json 结构", doc["sheet"] == "销售" and len(doc["rows"]) == 16
          and doc["rows"][1][3] == 990.0, str(doc["rows"][1] if "rows" in doc else doc))

    x1 = wpath("from_gbk.xlsx")
    expect_ok("convert gbk csv->xlsx", "convert", "-f",
              os.path.join(SAMPLES, "gbk.csv"), "--out", x1)
    out = expect_ok("read gbk 转换结果", "read", "-f", x1)
    check("convert gbk 中文+推断", out and out["rows"][1][0] == "张三"
          and out["rows"][1][2] == 1200.5 and out["rows"][1][3] is True,
          str(out and out["rows"][1]))

    x2 = wpath("from_json.xlsx")
    expect_ok("convert json->xlsx", "convert", "-f", o_json, "--out", x2)
    out = expect_ok("read json 转换结果", "read", "-f", x2, "--sheet", "Sheet1")
    check("convert json 回写", out and out["rows"][0][0] == "月份", str(out and out["rows"][0]))

    o_csv2 = wpath("out2.json")
    expect_ok("convert csv->json", "convert", "-f", os.path.join(SAMPLES, "utf8.csv"),
              "--out", o_csv2)

    # ---------- convert -> md / txt ----------
    msrc = wpath("md_src.xlsx")
    expect_ok("建 md/txt 源表", "write", "--create", "-f", msrc,
              "--data", json.dumps([["品名", "价格", "备注"],
                                     ["苹果|梨", 3.5, "好\n吃"],
                                     ["", True, ""]]))
    md_path = wpath("md_out.md")
    expect_ok("convert xlsx->md", "convert", "-f", msrc, "--out", md_path)
    with open(md_path, encoding="utf-8") as fh:
        md_text = fh.read()
    lines = md_text.splitlines()
    check("md 表头+分隔行", len(lines) >= 3
          and lines[0] == "| 品名 | 价格 | 备注 |"
          and lines[1] == "| --- | --- | --- |", md_text)
    check("md 转义 | 与换行", "苹果\\|梨" in md_text and "好<br>吃" in md_text
          and "| TRUE |" in md_text, md_text)

    txt_path = wpath("md_out.txt")
    expect_ok("convert xlsx->txt", "convert", "-f", msrc, "--out", txt_path)
    with open(txt_path, encoding="utf-8") as fh:
        txt_rows = list(csv.reader(io.StringIO(fh.read()), delimiter="\t"))
    check("txt TSV 无损回读", txt_rows == [["品名", "价格", "备注"],
                                          ["苹果|梨", "3.5", "好\n吃"],
                                          ["", "TRUE", ""]],
          repr(txt_rows))

    expect_ok("convert csv->md", "convert", "-f", os.path.join(SAMPLES, "utf8.csv"),
              "--out", wpath("utf8.md"))
    expect_ok("convert json->txt", "convert", "-f", o_json, "--out", wpath("out3.txt"))
    expect_err("convert xlsx->xls 仍不支持", "unsupported_format", "convert", "-f", demo,
               "--out", wpath("a.xls"))
    expect_err("convert 同路径", "same_file", "convert", "-f", demo, "--out", demo)

    # ---------- chart ----------
    ch = wpath("chart.xlsx")
    shutil.copy(cached, ch)
    out = expect_ok("chart add col", "chart", "add", "-f", ch, "--type", "col",
                    "--data", "A1:D13", "--at", "H2", "--title", "月度销售")
    check("chart series=3", out and out["series"] == 3 and out["anchor"] == "H2", str(out))
    expect_ok("chart add pie", "chart", "add", "-f", ch, "--type", "pie",
              "--data", "A1:B13")
    expect_err("chart pie 多系列报错", "bad_args", "chart", "add", "-f", ch,
               "--type", "pie", "--data", "A1:C13")
    expect_err("chart 单列报错", "bad_args", "chart", "add", "-f", ch,
               "--type", "col", "--data", "A1:A13")
    wb = load_workbook(ch)
    check("chart 保存后可重开", wb is not None, "")

    # ---------- image ----------
    im = wpath("image.xlsx")
    shutil.copy(demo, im)
    out = expect_ok("image add", "image", "add", "-f", im, "--image",
                    os.path.join(SAMPLES, "logo.png"), "--at", "B2", "--width", "120")
    check("image 尺寸", out and out["width_px"] == 120 and out["height_px"] == 48,
          str(out))
    expect_err("image 文件不存在", "no_file", "image", "add", "-f", im,
               "--image", wpath("no.png"), "--at", "B2")

    # ---------- pivot ----------
    out = expect_ok("pivot list", "pivot", "list", "-f", pivot_f)
    p0 = out["pivots"][0] if out and out["pivots"] else {}
    check("pivot list 内容", p0.get("name") == "产品透视"
          and p0.get("source_range") == "$A$1:$C$31" and p0.get("sheet") == "透视",
          json.dumps(p0, ensure_ascii=False))
    out = expect_ok("pivot set-source", "pivot", "set-source", "-f", pivot_f,
                    "--name", "产品透视", "--ref", "A1:C40")
    check("pivot set-source 范围", out and out["source_range"] == "$A$1:$C$40"
          and out["refresh_on_load"] is True, str(out))
    out = expect_ok("pivot list 验证", "pivot", "list", "-f", pivot_f)
    p0 = out["pivots"][0] if out and out["pivots"] else {}
    check("pivot 修改已落盘", p0.get("source_range") == "$A$1:$C$40",
          json.dumps(p0, ensure_ascii=False))
    expect_err("pivot 名字不存在", "no_pivot", "pivot", "set-source", "-f", pivot_f,
               "--name", "不存在", "--ref", "A1:B2")
    out = expect_ok("pivot list 无透视表文件", "pivot", "list", "-f", demo)
    check("pivot list 空", out and out["pivots"] == [], str(out))

    # ---------- layout(列宽/行高/冻结/筛选/隐藏) ----------
    lx = wpath("layout.xlsx")
    shutil.copy(demo, lx)
    out = expect_ok("layout 列宽行高", "layout", "-f", lx, "--sheet", "销售",
                    "--col-width", "A=18,B=20", "--row-height", "1=30")
    check("layout applied 2 项", out and len(out.get("applied", [])) == 2,
          str(out))
    lws = load_workbook(lx)["销售"]
    check("layout 列宽落盘", lws.column_dimensions["A"].width == 18
          and lws.row_dimensions[1].height == 30, "")
    out = expect_ok("layout 冻结+筛选", "layout", "-f", lx, "--sheet", "销售",
                    "--freeze", "A2", "--filter", "A1:E16")
    lws = load_workbook(lx)["销售"]
    check("layout 冻结/筛选落盘", lws.freeze_panes == "A2"
          and lws.auto_filter.ref == "A1:E16", str(lws.freeze_panes))
    expect_ok("layout 隐藏列 D:E", "layout", "-f", lx, "--sheet", "销售",
              "--hide-cols", "D:E")
    lws = load_workbook(lx)["销售"]
    check("layout 隐藏落盘", lws.column_dimensions["D"].hidden is True, "")
    expect_ok("layout 显示列", "layout", "-f", lx, "--sheet", "销售",
              "--show-cols", "D:E")
    expect_err("layout 非法列宽", "bad_args", "layout", "-f", lx,
               "--sheet", "销售", "--col-width", "Z=abc")

    # ---------- cond-format / comment / validate ----------
    cf = wpath("cf.xlsx")
    shutil.copy(demo, cf)
    expect_ok("cond-format gt", "cond-format", "add", "-f", cf,
              "--sheet", "销售", "--range", "B2:B16", "--op", "gt",
              "--value", "500")
    expect_ok("cond-format contains", "cond-format", "add", "-f", cf,
              "--sheet", "销售", "--range", "B2:B16", "--op", "contains",
              "--value", "月")
    expect_ok("cond-format duplicates", "cond-format", "add", "-f", cf,
              "--sheet", "销售", "--range", "A2:A16", "--op", "duplicates")
    cws = load_workbook(cf)["销售"]
    check("cond-format 落盘 2 区域", len(cws.conditional_formatting) == 2, "")
    expect_ok("cond-format clear --range", "cond-format", "clear", "-f", cf,
              "--sheet", "销售", "--range", "B2:B16")
    cws = load_workbook(cf)["销售"]
    check("cond-format 清除后剩 1", len(cws.conditional_formatting) == 1, "")
    expect_ok("cond-format clear --all", "cond-format", "clear", "-f", cf,
              "--sheet", "销售", "--all")
    cws = load_workbook(cf)["销售"]
    check("cond-format 全清", len(cws.conditional_formatting) == 0, "")
    code, out, err = run("cond-format", "clear", "-f", cf)
    check("cond-format 无 range 报错", code == 2, f"code={code}")

    cmf = wpath("cmt.xlsx")
    shutil.copy(demo, cmf)
    out = expect_ok("comment set", "comment", "set", "-f", cmf,
                    "--sheet", "销售", "--cell", "A2", "--text", "备注测试",
                    "--author", "tester")
    c2 = load_workbook(cmf)["销售"]["A2"].comment
    check("comment 落盘", c2 is not None and c2.text == "备注测试"
          and c2.author == "tester", str(c2))
    expect_ok("comment clear", "comment", "clear", "-f", cmf,
              "--sheet", "销售", "--cell", "A2")
    check("comment 已清除", load_workbook(cmf)["销售"]["A2"].comment is None, "")

    vf = wpath("val.xlsx")
    shutil.copy(demo, vf)
    out = expect_ok("validate add 列表", "validate", "add", "-f", vf,
                    "--sheet", "销售", "--range", "C2:C16", "--list", "北京,上海")
    vws = load_workbook(vf)["销售"]
    check("validate 落盘", out and vws.data_validations.dataValidation
          and vws.data_validations.dataValidation[0].formula1 == '"北京,上海"',
          str(out))
    expect_ok("validate add 区域源", "validate", "add", "-f", vf,
              "--sheet", "销售", "--range", "D2:D16", "--source", "人员!A1:A3")
    expect_ok("validate clear", "validate", "clear", "-f", vf,
              "--sheet", "销售", "--range", "C2:D16")
    check("validate 已清空", len(load_workbook(vf)["销售"]
                                 .data_validations.dataValidation) == 0, "")

    # ---------- insert / delete(行/列,合并区预检) ----------
    od = wpath("io.xlsx")
    shutil.copy(demo, od)
    expect_ok("io 先解除底部合并", "merge", "-f", od, "--sheet", "销售",
              "--range", "A16:B16", "--unmerge")
    ws_src = load_workbook(od)["销售"]
    v31 = ws_src["A3"].value
    expect_ok("insert rows 3x2", "insert", "-f", od, "--sheet", "销售",
              "--rows", "3", "--count", "2")
    ows = load_workbook(od)["销售"]
    check("insert 数据下移", ows["A5"].value == v31,
          "v31=%r A5=%r dump=%s" % (v31, ows["A5"].value,
            [[ows.cell(r, c).value for c in range(1, 4)]
             for r in range(1, 10)]))
    expect_err("insert 撞合并区拒绝", "layout_conflict", "insert", "-f",
               demo, "--sheet", "销售", "--rows", "16")
    expect_ok("delete rows 3-4", "delete", "-f", od, "--sheet", "销售",
              "--rows", "3-4")
    ows = load_workbook(od)["销售"]
    check("delete 数据恢复", ows["A3"].value == v31, str(ows["A3"].value))
    vc2 = ws_src["C2"].value
    expect_ok("insert cols B", "insert", "-f", od, "--sheet", "销售",
              "--cols", "B")
    ows = load_workbook(od)["销售"]
    check("insert 列右移", ows["D2"].value == vc2, str(ows["D2"].value))
    expect_ok("delete cols B", "delete", "-f", od, "--sheet", "销售",
              "--cols", "B")
    ows = load_workbook(od)["销售"]
    check("delete 列恢复", ows["C2"].value == vc2, str(ows["C2"].value))
    expect_err("delete 行格式错", "bad_args", "delete", "-f", od,
               "--sheet", "销售", "--rows", "2-x")

    # ---------- replace ----------
    rf = wpath("rep.xlsx")
    shutil.copy(demo, rf)
    expect_ok("replace 前置写文本", "write", "-f", rf, "--sheet", "销售",
              "--cell", "F1", "--data", '"Hello World ABC"')
    out = expect_ok("replace 基本", "replace", "-f", rf, "--sheet", "销售",
                    "--find", "world", "--replace", "X", "--range", "F1:F1")
    check("replace 大小写不敏感计数", out and out["cells_changed"] == 1
          and out["occurrences"] == 1, str(out))
    expect_ok("replace match-case", "replace", "-f", rf, "--sheet", "销售",
              "--find", "abc", "--replace", "Y", "--range", "F1:F1",
              "--match-case")
    out = expect_ok("replace 读取验证", "read", "-f", rf, "--sheet", "销售",
                    "--range", "F1:F1")
    check("replace 结果文本", out and out["rows"][0][0] == "Hello X ABC",
          str(out and out["rows"]))
    expect_ok("replace regex", "replace", "-f", rf, "--sheet", "销售",
              "--find", "^Hello", "--replace", "Hi", "--range", "F1:F1",
              "--regex")
    out = expect_ok("replace 结果=开头存文本", "replace", "-f", rf,
                    "--sheet", "销售", "--find", "Hi X", "--replace", "=NOW()",
                    "--range", "F1:F1")
    f1c = load_workbook(rf)["销售"]["F1"]
    check("replace =开头是文本", f1c.value == "=NOW() ABC"
          and f1c.data_type == "s", f"{f1c.value!r} {f1c.data_type}")
    expect_ok("replace 公式默认跳过", "replace", "-f", rf, "--sheet", "销售",
              "--find", "B2", "--replace", "Q9", "--range", "D2:D2")
    expect_ok("replace --in-formulas", "replace", "-f", rf, "--sheet", "销售",
              "--find", "B2", "--replace", "Q9", "--range", "D2:D2",
              "--in-formulas")
    d2 = load_workbook(rf)["销售"]["D2"]
    check("replace 公式内替换", d2.value == "=Q9*C2", str(d2.value))

    # ======================================================================
    # word 组(显式带组名;word write 增删改查 + 图片/样式元数据)
    # ======================================================================
    w1 = wpath("w_new.docx")
    out = expect_ok("word write 新建", "word", "write", "-f", w1, "--create",
                    "--text", "你好世界")
    check("word write 返回", out and out.get("appended_blocks") == 1, str(out))
    out = expect_ok("word write 追加", "word", "write", "-f", w1,
                    "--text", "第二段内容")
    out = expect_ok("word read 回读", "word", "read", "-f", w1)
    check("word 追加生效", out and len(out["paragraphs"]) == 2
          and out["paragraphs"][1]["text"] == "第二段内容",
          str(out and out["paragraphs"]))
    expect_err("word write 需 --create", "no_file", "word", "write",
               "-f", wpath("nope.docx"), "--text", "x")

    # ---------- word replace(模板占位符) ----------
    wt = wpath("wt.docx")
    expect_ok("word replace 模板准备", "word", "write", "-f", wt,
              "--create", "--text", "甲方 ${客户} 于 {日期} 签约")
    with open(wpath("repl.json"), "w", encoding="utf-8") as fh:
        json.dump({"${客户}": "测试公司", "{日期}": "2026-05-01"}, fh,
                  ensure_ascii=False)
    out = expect_ok("word replace 模板", "word", "replace", "-f", wt,
                    "--data-file", wpath("repl.json"))
    check("word replace 计数", out and out["occurrences"] == 2
          and out["rebuilt_paragraphs"] == 0, str(out))
    out = expect_ok("word replace 回读", "word", "read", "-f", wt)
    joined = "".join(p.get("text", "")
                      for p in (out or {}).get("paragraphs", []))
    check("word replace 内容", "甲方 测试公司 于 2026-05-01 签约" in joined,
          str(joined))
    expect_ok("word replace find 模式", "word", "replace", "-f", wt,
              "--find", "测试公司", "--replace", "新公司")
    expect_err("word replace 缺文件", "no_file", "word", "replace", "-f",
               wpath("nope.docx"), "--find", "a", "--replace", "b")

    with open(wpath("bad.json"), "w") as fh:
        fh.write("{not json")
    expect_err("word write 坏 json", "bad_json", "word", "write", "-f", w1,
               "--data-file", wpath("bad.json"))

    out = expect_ok("word read 样本结构", "word", "read", "-f", sample_docx)
    txts = [p["text"] for p in (out or {}).get("paragraphs", [])]
    check("word 样本标题", out and "Word 样本文档" in txts, str(out and txts))
    check("word 样本表格", out and len(out["tables"]) == 1
          and out["tables"][0]["cols"] == 2, str(out and out["tables"]))
    check("word 样本图片元数据", out and len(out["images"]) == 1
          and out["images"][0]["size_bytes"] > 0, str(out and out["images"]))

    imgdir = wpath("wimgs")
    out = expect_ok("word read --save-images", "word", "read", "-f",
                    sample_docx, "--save-images", imgdir)
    check("word 图片导出", out and out["saved_images"]
          and os.path.exists(out["saved_images"][0]), str(out))

    if os.path.exists(legacy_doc):
        out = expect_ok("word read legacy.doc 自动升级", "word", "read", "-f",
                        legacy_doc)
        check("doc 内容可读", out and len(out["paragraphs"]) > 0,
              str(out and len(out["paragraphs"])))
    if os.path.exists(legacy_xls):
        out = expect_ok("excel read legacy.xls 自动升级", "read", "-f",
                        legacy_xls)
        check("xls 内容可读", out and out["row_count"] >= 12, str(out))

    # ======================================================================
    # md 组(md → pdf/html/docx + 往返)
    # ======================================================================
    pdf_path = wpath("sample.pdf")
    out = expect_ok("md to-pdf", "md", "to-pdf", "-f", sample_md,
                    "--out", pdf_path)
    n_pdf = (out or {}).get("pages") or 1
    check("md to-pdf 页数", out and n_pdf >= 1 and n_pdf <= 3,
          str(out))
    html_path = wpath("sample.html")
    out = expect_ok("md to-html", "md", "to-html", "-f", sample_md,
                    "--out", html_path)
    with open(html_path, encoding="utf-8") as fh:
        html_txt = fh.read()
    check("html 含 mermaid 与样式", "mermaid" in html_txt
          and "codehilite" in html_txt and "@page" in html_txt, "")

    docx2 = wpath("from_md.docx")
    out = expect_ok("md to-docx", "md", "to-docx", "-f", sample_md,
                    "--out", docx2)
    out = expect_ok("md→docx→md 往返", "convert", "-f", docx2, "--to", "md",
                    "--out", wpath("roundtrip.md"))
    with open(wpath("roundtrip.md"), encoding="utf-8") as fh:
        rt = fh.read()
    check("往返保留标题", "# 办公文档转换样本" in rt, "")
    check("往返保留表格行", "| 苹果 |" in rt or "苹果" in rt, "")
    check("往返保留代码", "def calc" in rt, "")
    check("往返保留引用", ">" in rt and "引用块" in rt, "")

    # ======================================================================
    # pdf 组(基于 sample.pdf)
    # ======================================================================
    out = expect_ok("pdf info", "pdf", "info", "-f", pdf_path)
    fp = (out or {}).get("first_page") or {}
    check("pdf info 尺寸 A4", out and 209 <= fp.get("width_mm", 0) <= 212
          and 296 <= fp.get("height_mm", 0) <= 300, str(out))
    out = expect_ok("pdf read 文本", "pdf", "read", "-f", pdf_path)
    all_text = "".join(pg["text"] for pg in (out or {}).get("pages", []))
    check("pdf read 内容", out and "办公文档转换样本" in all_text, "")
    out = expect_ok("pdf read --pages 1", "pdf", "read", "-f", pdf_path,
                    "--pages", "1")
    check("pdf read 只取 1 页", out and len(out["pages"]) == 1, str(out))

    two_pdf = wpath("two.pdf")
    out = expect_ok("pdf merge", "pdf", "merge", "-f", pdf_path, pdf_path,
                    "--out", two_pdf)
    out = expect_ok("pdf merge 页数", "pdf", "info", "-f", two_pdf)
    check("merge 页数翻倍", out and out["pages"] == n_pdf * 2, str(out))
    code, _, _ = run("pdf", "merge", "-f", pdf_path, pdf_path)
    check("pdf merge 缺 --out 报错", code == 2, f"code={code}")

    split_dir = wpath("split")
    os.makedirs(split_dir, exist_ok=True)
    out = expect_ok("pdf split 全部", "pdf", "split", "-f", two_pdf,
                    "--out-dir", split_dir)
    files = [f for f in os.listdir(split_dir) if f.lower().endswith(".pdf")]
    check("split 文件数=页数", out and len(files) == n_pdf * 2, str(files))
    for f in files:
        os.remove(os.path.join(split_dir, f))
    out = expect_ok("pdf split 指定页", "pdf", "split", "-f", two_pdf,
                    "--pages", "1,2", "--out-dir", split_dir)
    files = [f for f in os.listdir(split_dir) if f.lower().endswith(".pdf")]
    check("split 指定页数", out and len(files) == 2, str(files))
    for f in files:
        os.remove(os.path.join(split_dir, f))

    rot = wpath("rot.pdf")
    expect_ok("pdf rotate", "pdf", "rotate", "-f", two_pdf, "--angle", "90",
              "--out", rot)
    out = expect_ok("pdf rotate 后 info", "pdf", "info", "-f", rot)
    rfp = (out or {}).get("first_page") or {}
    check("rotate 变横向", out and rfp["width_mm"] > rfp["height_mm"],
          str(rfp))

    enc = wpath("enc.pdf")
    expect_ok("pdf encrypt", "pdf", "encrypt", "-f", two_pdf,
              "--password", "pw123", "--out", enc)
    out = expect_ok("pdf info 加密文件", "pdf", "info", "-f", enc)
    check("encrypted 状态", out and out["encrypted"] is True
          and out["pages"] is None, str(out))
    expect_err("pdf read 加密文件", "encrypted", "pdf", "read", "-f", enc)
    expect_err("pdf encrypt 二次加密", "encrypted", "pdf", "encrypt", "-f",
               enc, "--password", "x", "--out", wpath("e2.pdf"))
    expect_err("pdf decrypt 错密码", "bad_password", "pdf", "decrypt", "-f",
               enc, "--password", "wrong", "--out", wpath("d1.pdf"))
    expect_ok("pdf decrypt", "pdf", "decrypt", "-f", enc,
              "--password", "pw123", "--out", wpath("d2.pdf"))
    expect_err("pdf decrypt 未加密", "not_encrypted", "pdf", "decrypt", "-f",
               two_pdf, "--password", "pw123", "--out", wpath("d3.pdf"))

    wm = wpath("wm.pdf")
    out = expect_ok("pdf watermark", "pdf", "watermark", "-f", pdf_path,
                    "--text", "内部资料", "--out", wm)
    check("watermark 返回", out and out["watermarked_pages"] == n_pdf, str(out))
    img_dir2 = wpath("pdimg")
    out = expect_ok("pdf to-image", "pdf", "to-image", "-f", pdf_path,
                    "--out-dir", img_dir2, "--dpi", "100")
    pngs = [f for f in os.listdir(img_dir2) if f.endswith(".png")]
    check("to-image 产物", out and len(pngs) == n_pdf, str(pngs))
    out = expect_ok("pdf images 提取", "pdf", "images", "-f", pdf_path,
                    "--out-dir", wpath("pdimgs"))
    check("pdf 内嵌图", out and out["count"] >= 1, str(out))
    # 无图 PDF: 纯文字 md 渲染(避免 two.pdf 内含 logo 误命中)
    plain_md = wpath("plain.md")
    with open(plain_md, "w", encoding="utf-8") as fh:
        fh.write("# 纯文本页\n\n没有任何图片的文档。\n")
    noimg_pdf = wpath("noimg.pdf")
    expect_ok("md to-pdf 纯文本", "md", "to-pdf", "-f", plain_md,
              "--out", noimg_pdf)
    expect_err("pdf images 无图报错", "no_images", "pdf", "images", "-f",
               noimg_pdf, "--out-dir", wpath("pdimgs2"))

    # ---------- pdf search / footer / from-images ----------
    out = expect_ok("pdf search 命中", "pdf", "search", "-f", pdf_path,
                    "--find", "办公文档转换样本")
    check("pdf search 结果", out and out["total"] >= 1
          and out["matches"][0]["count"] >= 1, str(out))
    out = expect_ok("pdf search 无命中", "pdf", "search", "-f", pdf_path,
                    "--find", "绝不存在的词XYZ")
    check("pdf search 空结果", out and out["total"] == 0, str(out))
    code, out, err = run("pdf", "search", "-f", pdf_path)
    check("pdf search 缺 --find", code == 2, f"code={code}")
    fpdf = wpath("foot.pdf")
    shutil.copy(pdf_path, fpdf)
    expect_ok("pdf footer 中文页码", "pdf", "footer", "-f", fpdf,
              "--text", "第 {page} 页 / 共 {pages} 页")
    out = expect_ok("pdf footer 后可搜", "pdf", "search", "-f", fpdf,
                    "--find", "第 1 页")
    check("pdf footer 文本写入", out and out["total"] >= 1, str(out))
    expect_ok("pdf footer 指定页", "pdf", "footer", "-f", fpdf,
              "--text", "机密", "--pages", "1")
    expect_err("pdf footer 页越界", "bad_args", "pdf", "footer", "-f",
               fpdf, "--text", "x", "--pages", "99-100")
    out = expect_ok("pdf from-images", "pdf", "from-images", "--out",
                    wpath("combo.pdf"), os.path.join(SAMPLES, "logo.png"),
                    os.path.join(SAMPLES, "logo.png"))
    check("pdf from-images 页数", out and out["pages"] == 2, str(out))
    out = expect_ok("pdf from-images info", "pdf", "info", "-f",
                    wpath("combo.pdf"))
    check("pdf from-images 可读", out and out["pages"] == 2, str(out))
    expect_err("pdf from-images 缺图", "no_file", "pdf", "from-images",
               "--out", wpath("combo2.pdf"), wpath("no.png"))

    # ======================================================================
    # convert 跨格式(WPS 引擎部分)
    # ======================================================================
    out = expect_ok("convert docx→pdf(WPS)", "convert", "-f", sample_docx,
                    "--out", wpath("w.pdf"))
    out = expect_ok("convert pdf→docx", "convert", "-f", two_pdf, "--to",
                    "docx", "--out", wpath("pdf2.docx"))
    check("pdf2docx 引擎", out and out.get("engine") == "pdf2docx", str(out))
    expect_ok("convert docx→md", "convert", "-f", sample_docx, "--to", "md",
              "--out", wpath("doc2.md"))
    with open(wpath("doc2.md"), encoding="utf-8") as fh:
        d2m_txt = fh.read()
    check("docx→md 标题", "Word 样本文档" in d2m_txt, "")
    expect_err("convert 未知扩展", "unsupported_format", "convert", "-f",
               sample_md, "--to", "xyz", "--out", wpath("a.xyz"))
    if os.path.exists(legacy_xls):
        out = expect_ok("convert xls→xlsx(WPS)", "convert", "-f", legacy_xls,
                        "--out", wpath("legacy_new.xlsx"))
        check("xls 转换警告", out and any("WPS" in w or "升级" in w
                                          for w in out.get("warnings", [])),
              str(out))
    if os.path.exists(legacy_doc):
        out = expect_ok("convert doc→pdf(WPS)", "convert", "-f", legacy_doc,
                        "--out", wpath("ld.pdf"))

    # ======================================================================
    # ppt 组(python-pptx 读写;WPS KWPP 相关需 legacy.ppt 样本存在)
    # ======================================================================
    out = expect_ok("ppt read 全量", "ppt", "read", "-f", sample_pptx)
    slides = (out or {}).get("slides") or []
    check("ppt read 页数", out and out["slides_total"] == 4
          and len(slides) == 4, str(out))
    s2 = slides[1] if len(slides) > 1 else {}
    check("ppt read 标题", s2.get("title") == "进展摘要", str(s2))
    t1 = (s2.get("texts") or [])[0]
    check("ppt read 文本", s2 and "需求评审完成" in str(t1), str(s2))
    check("ppt read 层级", any(isinstance(x, dict)
                                and x.get("level") == 1
                                and "覆盖 3 个部门" in x.get("text", "")
                                for x in (s2.get("texts") or [])), str(s2))
    check("ppt read 备注", s2.get("notes") == "强调时间点", str(s2))
    s3 = slides[2] if len(slides) > 2 else {}
    tab = (s3.get("tables") or [[]])[0]
    check("ppt read 表格", tab and tab[0] == ["指标", "Q1", "Q2"]
          and tab[2] == ["用户", "10", "22"], str(s3))
    s4 = slides[3] if len(slides) > 3 else {}
    check("ppt read 图片计数", s4 and s4.get("pictures", 0) >= 1, str(s4))
    out = expect_ok("ppt read 单页", "ppt", "read", "-f", sample_pptx,
                    "--slide", "2")
    check("ppt read 单页结构", out and out["pages"] == 1
          and out["slides"][0]["title"] == "进展摘要", str(out))
    expect_err("ppt read 页越界", "bad_args", "ppt", "read", "-f",
               sample_pptx, "--slide", "9")
    out = expect_ok("ppt read 导出图片", "ppt", "read", "-f", sample_pptx,
                    "--save-images", wpath("ppt_imgs"))
    check("ppt read 图片文件", out and len(out.get("images_saved") or []) >= 1
          and os.path.exists(os.path.join(wpath("ppt_imgs"),
                                          out["images_saved"][0])),
          str(out))

    def _ppt_json(name: str, obj) -> str:
        p = wpath(name)
        with open(p, "w", encoding="utf-8") as fh:
            json.dump(obj, fh, ensure_ascii=False)
        return p

    pj = _ppt_json("newppt.json", {"slides": [
        {"layout": "title", "title": "新建标题", "notes": "备注A"},
        {"layout": "title-content", "title": "要点页",
         "bullets": ["一", {"text": "一甲", "level": 1}]},
    ]})
    np_ = wpath("new.pptx")
    out = expect_ok("ppt write 新建", "ppt", "write", "-f", np_,
                    "--create", "--data-file", pj)
    check("ppt write 新建页数", out and out["slides_added"] == 2
          and out["created"] is True, str(out))
    out = expect_ok("ppt write 追加", "ppt", "write", "-f", np_,
                    "--data-file", pj)
    check("ppt write 追加页数", out and out["slides_added"] == 2
          and out["created"] is False, str(out))
    out = expect_ok("ppt write 后读回", "ppt", "read", "-f", np_)
    check("ppt write 总页数", out and out["slides_total"] == 4, str(out))
    chk = (out or {}).get("slides") or []
    check("ppt write 标题回读", chk and chk[0].get("title") == "新建标题"
          and chk[0].get("notes") == "备注A", str(chk and chk[0]))
    expect_err("ppt write 缺文件", "no_file", "ppt", "write", "-f",
               wpath("ghost.pptx"), "--data-file", pj)
    expect_err("ppt write 坏 JSON 结构", "bad_args", "ppt", "write",
               "-f", np_, "--data-file",
               _ppt_json("bad.json", {"slides": []}))
    if os.path.exists(legacy_ppt):
        out = expect_ok("ppt read legacy.ppt 升级", "ppt", "read", "-f",
                        legacy_ppt)
        check("legacy.ppt 页数", out and out["slides_total"] == 4, str(out))
        out = expect_ok("ppt to-pdf(WPS)", "ppt", "to-pdf", "-f",
                        sample_pptx, "--out", wpath("ppt_out.pdf"))
        check("ppt to-pdf 产物", out and out.get("pages")
              and os.path.exists(wpath("ppt_out.pdf")), str(out))
        expect_ok("ppt to-pdf 老格式", "ppt", "to-pdf", "-f", legacy_ppt,
                  "--out", wpath("ppt_old.pdf"))

    # ---------- rag prep ----------
    rsrc = wpath("rag_src.xlsx")
    expect_ok("rag 建源表", "write", "--create", "-f", rsrc,
              "--data", json.dumps([["月份", "销量"], ["2024-01", 990.0]]))
    out = expect_ok("rag prep 单文件", "rag", "prep", "-f", rsrc,
                    "--out-dir", WORK)
    check("rag prep 产物落盘", out and out["md"].startswith(WORK)
          and os.path.exists(out["md"]) and os.path.exists(out["json"]),
          str(out))
    with open(out["md"], encoding="utf-8") as fh:
        md_text = fh.read()
    check("rag md 语义保留", "月份" in md_text and "2024-01" in md_text
          and "990" in md_text)
    with open(out["json"], encoding="utf-8") as fh:
        doc = json.load(fh)
    sh0 = (doc.get("sheets") or [])[0]
    blk0 = (sh0.get("blocks") or [{}])[0]
    check("rag json 表头/行", doc.get("format") == "xlsx" and sh0
          and blk0.get("columns") == ["月份", "销量"]
          and (blk0.get("rows") or [])[0] == ["2024-01", 990.0],
          str(doc.get("sheets")))
    # demo.xlsx 含公式/合并区/多sheet → 公式全解 + 合并展开 + 行不丢
    out = expect_ok("rag prep 合并样本", "rag", "prep", "-f", demo,
                    "--out-dir", wpath("rag_dir"))
    with open(out["json"], encoding="utf-8") as fh:
        doc = json.load(fh)
    chk = (doc.get("sheets") or [])
    tbl_blocks = [b for s in chk for b in s.get("blocks", [])
                  if b.get("type") == "table"]
    total_rows = sum(len(b.get("rows") or []) for b in tbl_blocks)
    check("rag 合并展开行完整", tbl_blocks and total_rows >= 18
          and any(b.get("columns") and "月份" in b["columns"][0]
                  for b in tbl_blocks)
          and any("金额" in (b.get("columns") or []) for b in tbl_blocks),
          str(doc and total_rows))
    # 销售表公式无缓存(openpyxl 写死文本)→ 求值器全解, 合计行并入前表
    sales = next(s for s in chk if s["name"] == "销售")
    sb = sales["blocks"]
    unresolved = [u for b in sb if b["type"] == "table"
                  for u in (b.get("formula_unresolved") or [])]
    chk2 = (len(sb) == 1 and sb[0]["type"] == "table"
            and len(sb[0]["rows"]) == 14        # 12 数据行 + 2 合计行
            and not unresolved                      # 公式全部解出
            and sb[0]["row_numbers"][-1] == 16)    # 末尾原表行 16
    check("rag 公式全解+合计续接", chk2, str((sb, unresolved)))
    out = expect_ok("rag prep 目录批量", "rag", "prep", "-f", WORK,
                    "--out-dir", wpath("rag_dir"))
    check("rag qa.json 汇总", out and out.get("total", 0) >= 1
          and os.path.exists(wpath("rag_dir/qa.json")), str(out))
    expect_err("rag prep 不支持扩展", "unsupported_format", "rag", "prep",
               "-f", copy_sample("sample.md"))
    # rag_dir 里全是 .md/.json 产物 → 目录模式逐文件报 unsupported(不抛错)
    out = expect_ok("rag prep 目录全不支持逐条失败", "rag", "prep", "-f",
                    wpath("rag_dir"))
    check("rag 不支持逐条计 failed", out and out.get("total", 0) >= 3
          and out.get("failed") == out.get("total")
          and out.get("succeeded") == 0, str(out))
    empty_dir = wpath("rag_empty")
    os.makedirs(empty_dir, exist_ok=True)
    expect_err("rag prep 空目录", "no_file", "rag", "prep", "-f", empty_dir)
    # csv(标准库直写, 非 office-cli 自产)/txt/docx 链路
    import csv as _csv
    csvp = wpath("rag_src.csv")
    with open(csvp, "w", encoding="utf-8-sig", newline="") as fh:
        _csv.writer(fh).writerows([["型号", "数量"], ["A-1", 3], ["B-2", 5]])
    out = expect_ok("rag prep csv", "rag", "prep", "-f", csvp,
                    "--out-dir", WORK)
    with open(out["json"], encoding="utf-8") as fh:
        doc = json.load(fh)
    csvb = ((doc.get("sheets") or [{}])[0].get("blocks") or [])
    check("rag csv 表/行", doc.get("format") == "csv" and csvb
          and csvb[0]["type"] == "table"
          and csvb[0]["columns"] == ["型号", "数量"]
          and len(csvb[0]["rows"]) == 2, str(doc))
    txtp = wpath("rag_note.txt")
    with open(txtp, "w", encoding="utf-8") as fh:
        fh.write("交接班说明\n夜班共完成 12 批,全部合格。")
    out = expect_ok("rag prep txt", "rag", "prep", "-f", txtp,
                    "--out-dir", WORK)
    with open(out["json"], encoding="utf-8") as fh:
        doc = json.load(fh)
    txtb = ((doc.get("sheets") or [{}])[0].get("blocks") or [])
    check("rag txt 段落", doc.get("format") == "txt" and txtb
          and txtb[0]["type"] == "notes"
          and "夜班" in (txtb[0].get("text") or ""), str(doc))
    out = expect_ok("rag prep docx", "rag", "prep", "-f",
                    sample_docx, "--out-dir", wpath("rag_dir"))
    md_docx = open(out["md"], encoding="utf-8").read()
    check("rag docx 标题表保留", out and out["md"].startswith(WORK)
          and ("# " in md_docx or "| " in md_docx), str(out))
    # pptx: 每页 → sheet,要点 notes、表格 table(纯 python-pptx,不依赖 WPS)
    out = expect_ok("rag prep pptx", "rag", "prep", "-f", sample_pptx,
                    "--out-dir", WORK)
    with open(out["json"], encoding="utf-8") as fh:
        doc = json.load(fh)
    psh = doc.get("sheets") or []
    ptexts = [b.get("text") for s in psh for b in s.get("blocks", [])
              if b.get("type") == "notes"]
    ptabs = [b for s in psh for b in s.get("blocks", [])
             if b.get("type") == "table"]
    check("rag pptx 页/要点/表", doc.get("format") == "pptx" and len(psh) >= 4
          and any("需求评审完成" in (t or "") for t in ptexts)
          and ptabs and ptabs[0]["columns"][:2] == ["指标", "Q1"],
          str(doc)[:300])
    check("rag pptx md 层级", "## 第" in open(out["md"], encoding="utf-8").read())
    # rtf: WPS 升级为 docx 后同 docx 管道(手写最小 rtf,不依赖外部生成器)
    rtfp = wpath("rag_note.rtf")
    with open(rtfp, "w", encoding="ascii") as fh:
        body = "".join("\\u%d?" % ord(c) if ord(c) > 127 else c
                        for c in "交接班说明")
        body2 = "".join("\\u%d?" % ord(c) if ord(c) > 127 else c
                         for c in "夜班共完成 12 批,全部合格。")
        fh.write("{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0\\fnil\\fcharset134 SimSun;}}"
                 "\\f0\\fs24 " + body + "\\par\n" + body2 + "\\par\n}")
    out = expect_ok("rag prep rtf", "rag", "prep", "-f", rtfp,
                    "--out-dir", WORK)
    md_rtf = open(out["md"], encoding="utf-8").read()
    check("rag rtf 内容", out.get("format") == "rtf"
          and "夜班共完成 12 批" in md_rtf, md_rtf[:120])

    # rag prep --recalc: 数组公式等内置求值器算不出的公式靠 WPS 引擎兜底
    arr_x = wpath("rag_arr.xlsx")
    from openpyxl import Workbook as _WB
    _wb = _WB()
    _ws = _wb.active
    _ws.title = "明细"
    _ws["A1"], _ws["B1"], _ws["C1"] = "物料", "单价", "数量"
    for _i, (_n, _p, _q) in enumerate([("A", 10.0, 3), ("B", 20.0, 5)], start=2):
        _ws[f"A{_i}"], _ws[f"B{_i}"], _ws[f"C{_i}"] = _n, _p, _q
        _ws[f"D{_i}"] = f"=ROUND(B{_i}*C{_i},2)"
    _ws["F2"] = "=SUMPRODUCT((C2:C3>3)*B2:B3*C2:C3)"
    _wb.save(arr_x)
    out = expect_ok("rag prep 数组公式", "rag", "prep", "-f", arr_x,
                    "--out-dir", WORK)
    doc = json.load(open(out["json"], encoding="utf-8"))
    arr_unres = [u["formula"] for s in doc["sheets"] for b in s["blocks"]
                 for u in b.get("formula_unresolved") or []]
    check("rag 数组公式记 unresolved",
          any("SUMPRODUCT" in f for f in arr_unres), str(arr_unres))
    out = expect_ok("rag prep --recalc", "rag", "prep", "-f", arr_x,
                    "--out-dir", WORK, "--recalc")
    doc = json.load(open(out["json"], encoding="utf-8"))
    arr_unres = [u["formula"] for s in doc["sheets"] for b in s["blocks"]
                 for u in b.get("formula_unresolved") or []]
    arr_vals = [c for s in doc["sheets"] for b in s["blocks"]
                for row in b.get("rows") or [] for c in row]
    check("rag --recalc 算出数组公式", not arr_unres and any(
        isinstance(v, (int, float)) and abs(float(v) - 100.0) < 1e-6
        for v in arr_vals), f"{arr_unres} {str(arr_vals)[:160]}")
    check("rag --recalc 有来源说明",
          any("WPS" in w for w in doc.get("warnings") or []),
          str(doc.get("warnings")))

    # rag prep --engine formulas: 可选第三方后端(pip install 'office-cli[formula]')
    try:
        import importlib.util as _ilu
        _has_fm = _ilu.find_spec("formulas") is not None
    except (ImportError, ValueError):  # noqa: BLE001 未安装则跳过后端用例
        _has_fm = False
    if _has_fm:
        out = expect_ok("rag prep --engine formulas", "rag", "prep", "-f", arr_x,
                        "--out-dir", WORK, "--engine", "formulas")
        doc = json.load(open(out["json"], encoding="utf-8"))
        arr_unres = [u["formula"] for s in doc["sheets"] for b in s["blocks"]
                     for u in b.get("formula_unresolved") or []]
        arr_vals = [c for s in doc["sheets"] for b in s["blocks"]
                    for row in b.get("rows") or [] for c in row]
        check("rag --engine formulas 算出数组公式", not arr_unres and any(
            isinstance(v, (int, float)) and abs(float(v) - 100.0) < 1e-6
            for v in arr_vals), f"{arr_unres} {str(arr_vals)[:160]}")
        check("rag --engine formulas 有来源说明",
              any("formulas" in w for w in doc.get("warnings") or []),
              str(doc.get("warnings")))
    else:
        print("SKIP  rag prep --engine formulas(未安装 formulas 包)")
        expect_err("rag --engine formulas 缺依赖提示", "missing_dependency",
                   "rag", "prep", "-f", arr_x, "--out-dir", WORK,
                   "--engine", "formulas")

    # ---------- 引擎抽象(WPS / LibreOffice 双后端, 跨平台) ----------
    engine_checks()

    # ---------- rag 公式求值器(内存直测) ----------
    formula_checks()

    # ---------- rag 渲染层(number_format → 形态) ----------
    display_checks()

    # ---------- 汇总 ----------
    print()
    print(f"==== {_passed} passed, {_failed} failed ====")
    if _failures:
        print("failures:", *_failures, sep="\n  - ")
    return 1 if _failed else 0


def engine_checks() -> None:
    """引擎抽象层: WPS / LibreOffice 双后端、OFFICE_ENGINE 环境变量、CLI 参数优先级。"""
    out = expect_ok("info engines 字段", "info")
    eng = (out or {}).get("engines") or {}
    check("info 含 libreoffice 键", "libreoffice" in eng, str(eng))
    check("info 含 active 引擎名",
          eng.get("active") in ("WPS", "LibreOffice", "none"), str(eng))
    lo_path = eng.get("libreoffice")
    has_lo = bool(lo_path) and os.path.exists(lo_path)
    check("info libreoffice 路径有效",
          not lo_path or os.path.exists(lo_path), str(lo_path))

    # 带公式的小表,用于验证重算后端选择
    fx = wpath("eng_f.xlsx")
    from openpyxl import Workbook as _WB
    _wb = _WB()
    _w = _wb.active
    _w["A1"], _w["A2"], _w["A3"] = 1, 2, "=SUM(A1:A2)"
    _wb.save(fx)

    # OFFICE_ENGINE=none: 自动选引擎时明确报错(不静默降级)
    expect_err("OFFICE_ENGINE=none 拒绝重算", "engine_unavailable",
               "rag", "prep", "-f", fx, "--out-dir", WORK, "--recalc",
               env={"OFFICE_ENGINE": "none"})

    if not has_lo:
        print("SKIP  引擎用例(未安装 LibreOffice)")
        return

    # OFFICE_ENGINE=lo: 文档转 PDF 走 LibreOffice(跨平台路径)
    docx = copy_sample("sample.docx")
    out = expect_ok("OFFICE_ENGINE=lo docx→pdf", "convert", "-f", docx,
                    "--out", wpath("lo.pdf"), env={"OFFICE_ENGINE": "lo"})
    check("lo 转换产物存在", os.path.exists(wpath("lo.pdf")), "")
    check("lo 转换引擎标记",
          out and str(out.get("engine", "")).lower() in ("libreoffice", "lo"),
          str(out))

    # 旧格式升级也走 lo
    legacy_xls = os.path.join(SAMPLES, "legacy.xls")
    if os.path.exists(legacy_xls):
        out = expect_ok("OFFICE_ENGINE=lo xls→xlsx", "convert", "-f", legacy_xls,
                        "--out", wpath("lo_up.xlsx"), env={"OFFICE_ENGINE": "lo"})
        check("lo 升级警告带引擎名",
              out and any("LibreOffice" in w for w in out.get("warnings") or []),
              str(out.get("warnings") if out else out))

    # 公式重算: auto + OFFICE_ENGINE=lo → LibreOffice 重算并标明来源
    out = expect_ok("rag --recalc(OFFICE_ENGINE=lo)", "rag", "prep", "-f", fx,
                    "--out-dir", WORK, "--recalc", env={"OFFICE_ENGINE": "lo"})
    doc = json.load(open(out["json"], encoding="utf-8")) if out else {}
    check("lo 重算来源说明",
          any("LibreOffice" in w for w in doc.get("warnings") or []),
          str(doc.get("warnings")))

    # CLI 参数优先级高于环境变量: --engine lo 在 OFFICE_ENGINE=none 下仍可用
    out = expect_ok("--engine lo 覆盖 OFFICE_ENGINE=none", "rag", "prep", "-f", fx,
                    "--out-dir", WORK, "--engine", "lo",
                    env={"OFFICE_ENGINE": "none"})
    doc = json.load(open(out["json"], encoding="utf-8")) if out else {}
    check("--engine lo 来源说明",
          any("LibreOffice" in w for w in doc.get("warnings") or []),
          str(doc.get("warnings")))

    # ---------- 默认打印机隔离(引擎启动不该去连网络打印机) ----------
    code = (
        "import json\n"
        "from office import printer as P\n"
        "orig = P.current_default()\n"
        "with P.isolate() as iso:\n"
        "    inside = P.current_default()\n"
        "after = P.current_default()\n"
        "print(json.dumps({'orig': orig, 'inside': inside, 'after': after,\n"
        "                  'net': P.needs_isolation(), 'target': P.isolate_target(),\n"
        "                  'switched': iso.switched}))\n"
    )
    got = py_json(code)
    check("printer 隔离状态可读", isinstance(got, dict) and "orig" in got, str(got))
    if isinstance(got, dict):
        check("printer 退出后默认值还原", got["after"] == got["orig"],
              f"{got['orig']} -> {got['after']}")
        if got.get("net") and got.get("target"):
            check("网络默认打印机被临时切换",
                  got["switched"] is True and got["inside"] == got["target"], str(got))
        else:
            check("非网络默认打印机不动它",
                  got["switched"] is False and got["inside"] == got["orig"], str(got))

    # 引擎跑完不留副作用(默认打印机与转换前一致)
    _probe = ("import json; from office import printer as P; "
              "print(json.dumps(P.current_default()))")
    before = py_json(_probe)
    expect_ok("隔离下引擎转换", "convert", "-f", docx,
              "--out", wpath("iso.pdf"),
              env={"OFFICE_ENGINE": "lo" if has_lo else "wps"})
    after = py_json(_probe)
    check("引擎转换后默认打印机不变", before == after, f"{before} -> {after}")

    out = expect_ok("info 含 printer 段", "info")
    check("info engines.printer.default 可读",
          isinstance(((out or {}).get("engines") or {}).get("printer"), dict),
          str(((out or {}).get("engines") or {}).get("printer")))


if __name__ == "__main__":
    sys.exit(main())
